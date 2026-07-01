from __future__ import annotations

import io

from fastapi import APIRouter, Depends, File, Query, UploadFile
from fastapi.responses import StreamingResponse
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.api.deps import get_db_session, get_platform_db_session, require_platform_superadmin
from app.models.people import (
    EmployeeExt,
    EmployeeWorkInfo,
    LicenseAssignment,
    LicenseContract,
    OrgEmployee,
    OrgGroup,
    PeripheralInventory,
    SystemInventory,
)
from app.models.platform_person import DimPerson
from app.schemas.assets import ReconciliationResponse
from app.schemas.user import UserContext
from app.services import assets
from app.services.people_master_import import apply_people_master_import, build_people_master_plan

router = APIRouter(prefix="/ppl/bulk", tags=["bulk"])


@router.get("/reconciliation", response_model=ReconciliationResponse)
async def reconciliation(
    user: UserContext = Depends(require_platform_superadmin),
    db: AsyncSession = Depends(get_db_session),
    platform: AsyncSession = Depends(get_platform_db_session),
) -> ReconciliationResponse:
    return await assets.reconciliation_report(db, platform)


@router.post("/org-upload/preview")
async def preview_people_master_org_upload(
    file: UploadFile = File(...),
    user: UserContext = Depends(require_platform_superadmin),
    db: AsyncSession = Depends(get_db_session),
    platform: AsyncSession = Depends(get_platform_db_session),
) -> dict:
    contents = await file.read()
    plan, _rows = await build_people_master_plan(db, platform, contents, file.filename or "people-master.xlsx")
    return {"applied": False, "requires_override": bool(plan.issues and not plan.blocking_count), **plan.as_dict()}


@router.post("/org-upload/apply")
async def apply_people_master_org_upload(
    force: bool = Query(False),
    file: UploadFile = File(...),
    user: UserContext = Depends(require_platform_superadmin),
    db: AsyncSession = Depends(get_db_session),
    platform: AsyncSession = Depends(get_platform_db_session),
) -> dict:
    contents = await file.read()
    return await apply_people_master_import(
        db,
        platform,
        contents,
        file.filename or "people-master.xlsx",
        performed_by=user.person_id_platform or user.email,
        force=force,
    )


@router.get("/export/current")
async def export_current_data(
    user: UserContext = Depends(require_platform_superadmin),
    db: AsyncSession = Depends(get_db_session),
    platform: AsyncSession = Depends(get_platform_db_session),
) -> StreamingResponse:
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)

    people = (await platform.execute(select(DimPerson).order_by(DimPerson.person_code.asc()))).scalars().all()
    ext_rows = (await db.execute(select(EmployeeExt))).scalars().all()
    work_rows = (
        await db.execute(select(EmployeeExt.employee_number, EmployeeWorkInfo).join(EmployeeWorkInfo, EmployeeWorkInfo.employee_id == EmployeeExt.id))
    ).all()
    ext_by_emp = {row.employee_number: row for row in ext_rows if row.employee_number}
    work_by_emp = {emp_no: work for emp_no, work in work_rows}

    ws = wb.create_sheet("People")
    ws.append([
        "person_id", "employee_number", "name", "email", "platform_status",
        "employment_status", "worker_type", "time_type", "department", "sub_department",
        "business_unit", "job_title", "reporting_manager_id", "date_joined", "exit_date",
    ])
    for person in people:
        emp_no = str(person.person_code or "")
        ext = ext_by_emp.get(emp_no)
        work = work_by_emp.get(emp_no)
        ws.append([
            str(person.person_id or ""),
            emp_no,
            person.full_name or person.display_name or "",
            person.email or "",
            person.status or "",
            ext.employment_status if ext else "",
            ext.worker_type if ext else "",
            ext.time_type if ext else "",
            work.department if work else "",
            work.sub_department if work else "",
            work.business_unit if work else "",
            work.job_title if work else "",
            work.reporting_manager_id if work else "",
            work.date_joined.isoformat() if work and work.date_joined else "",
            work.exit_date.isoformat() if work and work.exit_date else "",
        ])

    groups = (await db.execute(select(OrgGroup).order_by(OrgGroup.principal_name, OrgGroup.sort_order))).scalars().all()
    ws = wb.create_sheet("Groups")
    ws.append(["group_key", "name", "principal_name", "team_lead_emp", "parent_name", "color_hex", "sort_order", "is_active"])
    for row in groups:
        ws.append([row.group_key, row.name, row.principal_name, row.team_lead_emp or "", row.parent_name or "", row.color_hex or "", row.sort_order, row.is_active])

    org_rows = (await db.execute(select(OrgEmployee).order_by(OrgEmployee.principal_name, OrgEmployee.group_key, OrgEmployee.employee_no))).scalars().all()
    people_by_emp = {str(p.person_code): p for p in people if p.person_code}
    group_by_key = {g.group_key: g for g in groups if g.group_key}
    ws = wb.create_sheet("Org")
    ws.append([
        "employee_no", "name", "email", "principal_name", "group_key", "group_name", "org_level",
        "include_in_org", "source_manager_emp", "manager_override_emp", "designation_level",
        "prior_exp_years", "sl_exp_years", "notes",
    ])
    for row in org_rows:
        person = people_by_emp.get(row.employee_no)
        group = group_by_key.get(row.group_key)
        ws.append([
            row.employee_no,
            (person.full_name or person.display_name) if person else "",
            person.email if person else "",
            row.principal_name,
            row.group_key,
            group.name if group else "",
            row.org_level,
            row.include_in_org,
            row.source_manager_emp or "",
            row.manager_override_emp or "",
            row.designation_level or "",
            float(row.prior_exp_years) if row.prior_exp_years is not None else "",
            float(row.sl_exp_years) if row.sl_exp_years is not None else "",
            row.notes or "",
        ])

    assignments = (await db.execute(select(LicenseAssignment).order_by(LicenseAssignment.work_email, LicenseAssignment.tool_name))).scalars().all()
    ws = wb.create_sheet("Licenses")
    ws.append(["work_email", "holder_name", "tool_name", "tool_short_name", "plan", "status", "assigned_on", "renewal_date", "notes"])
    people_by_email = {str(p.email or "").strip().lower(): p for p in people if p.email}
    for row in assignments:
        person = people_by_email.get(str(row.work_email or "").strip().lower())
        ws.append([
            row.work_email,
            (person.full_name or person.display_name) if person else "",
            row.tool_name,
            row.tool_short_name or "",
            row.plan or "",
            row.status,
            row.assigned_on.isoformat() if row.assigned_on else "",
            row.renewal_date.isoformat() if row.renewal_date else "",
            row.notes or "",
        ])

    contracts = (await db.execute(select(LicenseContract).order_by(LicenseContract.software, LicenseContract.end_date))).scalars().all()
    ws = wb.create_sheet("Contracts")
    ws.append(["contract_key", "entity", "software", "short_name", "category", "contract_no", "contract_type", "serial_no", "seats", "vendor", "start_date", "end_date", "cost", "currency", "status", "user_type", "notes"])
    for row in contracts:
        ws.append([
            row.contract_key, row.entity or "", row.software, row.short_name or "", row.category or "",
            row.contract_no or "", row.contract_type or "", row.serial_no or "", row.seats or 0,
            row.vendor or "", row.start_date.isoformat() if row.start_date else "", row.end_date.isoformat() if row.end_date else "",
            float(row.cost) if row.cost is not None else "", row.currency or "", row.status or "", row.user_type or "", row.notes or "",
        ])

    systems = (await db.execute(select(SystemInventory).order_by(SystemInventory.system_id))).scalars().all()
    ws = wb.create_sheet("Systems")
    ws.append([
        "system_id", "system_type", "assigned_email", "assigned_name", "stored_team", "derived_org_group",
        "processor", "ram_gb", "graphics_card", "storage", "os", "autocad_version", "sketchup_version",
        "3dsmax_version", "rhino_version", "enscape_version", "d5_render", "adobe_versions", "office_version",
        "antivirus", "purchase_date", "vendor", "service_tag", "serial_no", "capability_tier", "composite_score", "status", "notes",
    ])
    org_by_emp = {row.employee_no: row for row in org_rows if row.employee_no}
    for row in systems:
        person = people_by_email.get(str(row.assigned_email or "").strip().lower())
        org = org_by_emp.get(str(person.person_code)) if person and person.person_code else None
        group = group_by_key.get(org.group_key) if org else None
        ws.append([
            row.system_id, row.system_type or "", row.assigned_email or "",
            (person.full_name or person.display_name) if person else row.user_display or "",
            row.team or "", group.name if group else "", row.processor or "",
            float(row.ram_gb) if row.ram_gb is not None else "", row.graphics_card or "", row.storage or "",
            row.os or "", row.autocad_version or "", row.sketchup_version or "", row.threedmax_version or "",
            row.rhino_version or "", row.enscape_version or "", row.d5_render or "", row.adobe_versions or "",
            row.office_version or "", row.antivirus or "", row.purchase_date.isoformat() if row.purchase_date else "",
            row.vendor or "", row.service_tag or "", row.serial_no or "", row.capability_tier or "",
            row.composite_score or "", row.status, row.notes or "",
        ])

    peripherals = (await db.execute(select(PeripheralInventory).order_by(PeripheralInventory.category, PeripheralInventory.item_id))).scalars().all()
    ws = wb.create_sheet("Peripherals")
    ws.append(["item_id", "category", "item", "model", "serial", "quantity", "condition", "location", "assigned_to", "status", "notes"])
    for row in peripherals:
        ws.append([row.item_id, row.category or "", row.item, row.model or "", row.serial or "", row.quantity, row.condition or "", row.location or "", row.assigned_to or "", row.status, row.notes or ""])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="people-org-current-data.xlsx"'},
    )
