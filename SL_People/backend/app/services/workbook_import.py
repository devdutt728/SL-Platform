from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any, Optional

from openpyxl import load_workbook
from sqlalchemy import delete, func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.people import (
    EmployeeExt,
    EmployeeWorkInfo,
    LicenseAssignment,
    LicenseContract,
    OrgEmployee,
    OrgGroup,
    OrgPrincipal,
    PeripheralInventory,
    SystemInventory,
)
from app.models.platform_person import DimPerson
from app.services.console_logic import ORG_PRINCIPALS, compute_experience, grade_pc, level_for, short_name_for


@dataclass
class WorkbookImportResult:
    employees: int = 0
    groups: int = 0
    org_employees: int = 0
    license_assignments: int = 0
    license_contracts: int = 0
    systems: int = 0
    peripherals: int = 0
    warnings: list[str] = field(default_factory=list)


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _none(value: Any) -> Optional[str]:
    text = _text(value)
    if not text or text.lower() == "nan":
        return None
    return text


def _email(value: Any) -> Optional[str]:
    text = _none(value)
    return text.lower() if text else None


def _bool(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    text = _text(value).lower()
    return text in {"1", "true", "yes", "y", "active", "included"}


def _int(value: Any, default: int = 0) -> int:
    try:
        if value is None or value == "":
            return default
        return int(float(value))
    except (TypeError, ValueError):
        return default


def _float(value: Any) -> Optional[float]:
    try:
        if value is None or value == "":
            return None
        return float(value)
    except (TypeError, ValueError):
        return None


def _date(value: Any) -> Optional[date]:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return None


def _norm_assignment_status(value: Any) -> str:
    text = _text(value)
    lowered = text.lower()
    if lowered in {"license assigned", "assigned", "active"}:
        return "Assigned"
    if lowered in {"removed", "revoked"}:
        return "Removed"
    if lowered in {"unassigned", "available"}:
        return "Unassigned"
    return text or "Assigned"


def _worker_type(value: Any) -> str:
    return (_text(value) or "Permanent").lower().replace(" ", "_")


def _rows(path: Path, sheet_name: str) -> list[dict[str, Any]]:
    wb = load_workbook(path, read_only=False, data_only=True)
    ws = wb[sheet_name]
    headers = [_text(ws.cell(1, col).value) for col in range(1, ws.max_column + 1)]
    out: list[dict[str, Any]] = []
    for row_idx in range(2, ws.max_row + 1):
        values = {headers[col - 1]: ws.cell(row_idx, col).value for col in range(1, ws.max_column + 1) if headers[col - 1]}
        if any(value not in (None, "") for value in values.values()):
            out.append(values)
    return out


async def _platform_people(platform_session: AsyncSession) -> dict[str, DimPerson]:
    rows = (await platform_session.execute(select(DimPerson))).scalars().all()
    return {row.person_code: row for row in rows if row.person_code}


async def _seed_principals(session: AsyncSession) -> None:
    existing = {
        name for name in (await session.execute(select(OrgPrincipal.name))).scalars().all()
    }
    for order, principal in enumerate(ORG_PRINCIPALS, start=1):
        if principal["name"] not in existing:
            session.add(OrgPrincipal(name=principal["name"], color=principal["color"], sort_order=order))


async def import_org_workbook(
    session: AsyncSession,
    platform_session: AsyncSession,
    workbook_path: str | Path,
    *,
    performed_by: Optional[str] = None,
) -> WorkbookImportResult:
    path = Path(workbook_path)
    result = WorkbookImportResult()
    people_by_code = await _platform_people(platform_session)

    await _seed_principals(session)
    await session.flush()

    group_rows = _rows(path, "Groups")
    existing_groups = {
        row.group_key: row
        for row in (await session.execute(select(OrgGroup))).scalars().all()
    }
    for row in group_rows:
        key = _text(row.get("Group Key"))
        if not key or not _bool(row.get("Active")):
            continue
        group = existing_groups.get(key) or OrgGroup(group_key=key, name="")
        group.name = _text(row.get("Group / Team")) or key
        group.principal_name = _text(row.get("Principal"))
        group.parent_name = _none(row.get("Parent"))
        group.color_hex = _none(row.get("Color Hex"))
        group.sort_order = _int(row.get("Sort"), 999)
        group.is_active = True
        existing_groups[key] = group
        session.add(group)
        result.groups += 1

    master_rows = _rows(path, "Master_Employees")
    included_master = [
        row for row in master_rows
        if _text(row.get("Employee No")) and _bool(row.get("Include in Org"))
    ]
    lead_by_group: dict[str, str] = {}
    name_to_emp: dict[str, str] = {}
    for row in included_master:
        emp_no = _text(row.get("Employee No"))
        name = _text(row.get("Employee Name"))
        if name:
            name_to_emp[name.lower()] = emp_no
    for row in included_master:
        group_key = _text(row.get("Group Key"))
        lead_name = _text(row.get("Team Lead"))
        lead_emp = name_to_emp.get(lead_name.lower()) if lead_name else None
        if group_key and lead_emp:
            lead_by_group[group_key] = lead_emp
    for key, lead_emp in lead_by_group.items():
        if key in existing_groups:
            existing_groups[key].team_lead_emp = lead_emp

    existing_ext = {
        row.employee_number: row
        for row in (await session.execute(select(EmployeeExt))).scalars().all()
    }
    existing_work = {
        row.employee_id: row
        for row in (await session.execute(select(EmployeeWorkInfo))).scalars().all()
    }
    emp_id_by_no: dict[str, str] = {}
    for row in included_master:
        emp_no = _text(row.get("Employee No"))
        dp = people_by_code.get(emp_no)
        ext = existing_ext.get(emp_no) or EmployeeExt(employee_number=emp_no)
        ext.person_id = dp.person_id if dp else None
        ext.worker_type = _worker_type(row.get("Worker Type"))
        ext.employment_status = "working"
        ext.time_type = "fulltime"
        ext.is_deleted = False
        ext.updated_by_person_id = performed_by
        if not ext.created_by_person_id:
            ext.created_by_person_id = performed_by
        session.add(ext)
        await session.flush()

        manager_emp = _text(row.get("Manager Override Emp No")) or _text(row.get("Source Manager Emp No"))
        emp_id_by_no[emp_no] = ext.id
        work = existing_work.get(ext.id) or EmployeeWorkInfo(employee_id=ext.id)
        work.business_unit = _none(row.get("Business Unit"))
        work.department = _none(row.get("Department"))
        work.job_title = _none(row.get("Job Title"))
        work.date_joined = _date(row.get("DOJ"))
        work.reporting_manager_id = existing_ext.get(manager_emp).id if manager_emp in existing_ext else None
        session.add(work)
        result.employees += 1

    await session.flush()
    existing_ext = {
        row.employee_number: row
        for row in (await session.execute(select(EmployeeExt))).scalars().all()
    }
    for row in included_master:
        emp_no = _text(row.get("Employee No"))
        manager_emp = _text(row.get("Manager Override Emp No")) or _text(row.get("Source Manager Emp No"))
        ext = existing_ext.get(emp_no)
        if ext:
            work = await session.get(EmployeeWorkInfo, ext.id)
            if work and manager_emp in existing_ext:
                work.reporting_manager_id = existing_ext[manager_emp].id

    existing_org = {
        row.employee_no: row
        for row in (await session.execute(select(OrgEmployee))).scalars().all()
    }
    for row in included_master:
        emp_no = _text(row.get("Employee No"))
        group_key = _text(row.get("Group Key"))
        group = existing_groups.get(group_key)
        if not group:
            result.warnings.append(f"Skipping {emp_no}: unknown group {group_key}")
            continue
        title = _none(row.get("Job Title"))
        level = level_for(title)
        exp = compute_experience(_date(row.get("DOJ")), _float(row.get("Prior Exp (yrs)")))
        org = existing_org.get(emp_no) or OrgEmployee(employee_no=emp_no, group_key=group_key, principal_name=group.principal_name, org_level="Member")
        org.group_key = group_key
        org.principal_name = group.principal_name
        org.org_level = _text(row.get("Org Level")) or "Member"
        org.include_in_org = True
        org.source_manager_emp = _none(row.get("Source Manager Emp No"))
        org.manager_override_emp = _none(row.get("Manager Override Emp No"))
        org.designation_level = level.label
        org.designation_color = level.color
        org.designation_order = level.order
        org.prior_exp_years = _float(row.get("Prior Exp (yrs)"))
        org.sl_exp_years = exp.sl_exp_years
        org.o_exp_years = exp.o_exp_years
        org.image_url = _none(row.get("Image URL"))
        org.notes = _none(row.get("Notes"))
        org.updated_at = datetime.utcnow()
        org.updated_by_person_id = performed_by
        session.add(org)
        result.org_employees += 1

    await session.execute(delete(LicenseAssignment))
    for row in _rows(path, "Licenses"):
        email = _email(row.get("Work Email"))
        tool = _none(row.get("License / Tool"))
        if not email or not tool or "unassigned" in email:
            continue
        status = _norm_assignment_status(row.get("Status"))
        if status == "Removed":
            continue
        session.add(
            LicenseAssignment(
                work_email=email,
                tool_name=tool,
                tool_short_name=short_name_for(tool),
                plan=_none(row.get("Plan / People")) or _none(row.get("Plan / Seat")),
                status=status,
                assigned_on=_date(row.get("Assigned On")),
                renewal_date=_date(row.get("Renewal / End Date")),
                cost_centre=_none(row.get("Cost Centre / Team")),
                notes=_none(row.get("Notes")),
                smart_key=_none(row.get("Smart License Key")),
            )
        )
        result.license_assignments += 1

    existing_contracts = {
        row.contract_key: row
        for row in (await session.execute(select(LicenseContract))).scalars().all()
    }
    for row in _rows(path, "License_Inventory"):
        key = _text(row.get("Contract Key"))
        software = _none(row.get("Software"))
        if not key or not software:
            continue
        contract = existing_contracts.get(key) or LicenseContract(contract_key=key, software=software)
        existing_contracts[key] = contract
        contract.entity = _none(row.get("Entity"))
        contract.software = software
        contract.short_name = short_name_for(software)
        contract.category = _none(row.get("Category"))
        contract.contract_no = _none(row.get("Contract No"))
        contract.contract_type = _none(row.get("Contract Type"))
        contract.serial_no = _none(row.get("Serial No"))
        contract.seats = _int(row.get("People"))
        contract.start_date = _date(row.get("Start Date"))
        contract.end_date = _date(row.get("End Date"))
        contract.vendor = _none(row.get("Vendor"))
        contract.status = _none(row.get("Status"))
        contract.user_type = _none(row.get("User Type"))
        contract.notes = _none(row.get("Notes"))
        contract.updated_at = datetime.utcnow()
        session.add(contract)
        result.license_contracts += 1

    existing_systems = {
        row.system_id: row
        for row in (await session.execute(select(SystemInventory))).scalars().all()
    }
    master_by_email = {
        _email(row.get("Work Email")): row
        for row in included_master
        if _email(row.get("Work Email"))
    }
    for row in _rows(path, "Systems"):
        system_id = _text(row.get("System ID"))
        if not system_id or "👉" in system_id:
            continue
        email = _email(row.get("Assigned Email"))
        master = master_by_email.get(email)
        system = existing_systems.get(system_id) or SystemInventory(system_id=system_id)
        existing_systems[system_id] = system
        system.system_type = _none(row.get("Type"))
        system.assigned_email = email
        system.user_display = _none(master.get("Employee Name")) if master else _none(row.get("User Display"))
        system.team = _none(master.get("Group / Team")) if master else _none(row.get("Team"))
        system.processor = _none(row.get("Processor"))
        system.ram_gb = _float(row.get("RAM (GB)"))
        system.ram_slots_free = _none(row.get("RAM Slots Free"))
        system.graphics_card = _none(row.get("Graphics Card"))
        system.cpu_cores = _none(row.get("CPU Cores"))
        system.storage = _none(row.get("Storage"))
        system.motherboard = _none(row.get("Motherboard"))
        system.os = _none(row.get("OS"))
        system.office_version = _none(row.get("Office Version"))
        system.autocad_version = _none(row.get("AutoCAD Version"))
        system.adobe_versions = _none(row.get("Adobe Versions"))
        system.sketchup_version = _none(row.get("SketchUp Version"))
        system.threedmax_version = _none(row.get("3ds Max Version"))
        system.rhino_version = _none(row.get("Rhino Version"))
        system.enscape_version = _none(row.get("Enscape Version"))
        system.d5_render = _none(row.get("D5 Render"))
        system.antivirus = _none(row.get("Antivirus"))
        system.purchase_date = _date(row.get("Purchase Date"))
        system.vendor = _none(row.get("Vendor"))
        system.service_tag = _none(row.get("Service Tag"))
        system.serial_no = _none(row.get("Serial No"))
        grade = grade_pc(system.processor, system.graphics_card, str(system.ram_gb or ""))
        system.composite_score = grade.score
        system.capability_tier = grade.tier
        system.upgrade_suggestion = grade.suggestion
        system.status = _none(row.get("Status")) or "Active"
        system.notes = _none(row.get("Notes"))
        system.updated_at = datetime.utcnow()
        session.add(system)
        result.systems += 1

    existing_peripherals = {
        row.item_id: row
        for row in (await session.execute(select(PeripheralInventory))).scalars().all()
    }
    for row in _rows(path, "Peripherals"):
        item_id = _text(row.get("Item ID"))
        item_name = _none(row.get("Item"))
        if not item_id or "👉" in item_id or not item_name:
            continue
        item = existing_peripherals.get(item_id) or PeripheralInventory(item_id=item_id, item=item_name)
        existing_peripherals[item_id] = item
        item.category = _none(row.get("Category"))
        item.item = item_name
        item.model = _none(row.get("Model"))
        item.serial = _none(row.get("Serial"))
        item.quantity = _int(row.get("Quantity"), 1)
        item.condition = _none(row.get("Condition"))
        item.location = _none(row.get("Location"))
        item.assigned_to = _none(row.get("Assigned To"))
        item.status = _none(row.get("Status")) or "Active"
        item.notes = _none(row.get("Notes"))
        item.updated_at = datetime.utcnow()
        session.add(item)
        result.peripherals += 1

    await session.commit()
    return result
