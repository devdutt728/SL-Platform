"""Import AutoCAD LT contracts and users from Autodesk exports.

The users workbook is expected to contain a `Users` sheet with columns:
team_alias, first_name, last_name, email, autodesk_id, role, account_status,
offering_name, product_name.

Usage:
    python -m migrations.0030_import_autodesk_autocad --file "C:/Users/Dev/Downloads/Autocad.xlsx"
    python -m migrations.0030_import_autodesk_autocad --file "..." --apply
"""

from __future__ import annotations

import argparse
import asyncio
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import select

from app.db.session import SessionLocal, engine
from app.models.people import LicenseAssignment, LicenseContract
from app.services.console_logic import short_name_for


AUTOCAD_TOOL = "AutoCAD LT"

AUTOCAD_CONTRACTS = [
    ("110003375737", 10, date(2026, 12, 22)),
    ("110003617657", 10, date(2026, 6, 22)),
    ("110003653231", 6, date(2026, 7, 18)),
    ("110003696736", 10, date(2026, 8, 29)),
    ("110004469439", 10, date(2026, 7, 26)),
    ("110004850190", 10, date(2026, 8, 29)),
    ("110004850159", 20, date(2026, 8, 29)),
    ("110004877425", 15, date(2026, 10, 3)),
    ("110005089936", 25, date(2026, 11, 9)),
]


@dataclass
class AutodeskImportStats:
    autocad_users: int = 0
    skipped_rows: int = 0
    deleted_assignments: int = 0
    created_assignments: int = 0
    deleted_contracts: int = 0
    created_contracts: int = 0
    active_contract_seats: int = 0
    expired_contract_seats: int = 0


def _text(value: Any) -> str:
    return str(value or "").strip()


def _email(value: Any) -> str:
    return _text(value).lower()


def _load_autocad_users(path: Path) -> tuple[list[dict[str, Any]], int]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["Users"]
    headers = [_text(value) for value in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
    users_by_email: dict[str, dict[str, Any]] = {}
    skipped = 0
    for values in ws.iter_rows(min_row=2, values_only=True):
        row = dict(zip(headers, values))
        email = _email(row.get("email"))
        product = _text(row.get("product_name") or row.get("offering_name"))
        status = _text(row.get("account_status")).lower()
        if not email or short_name_for(product) != AUTOCAD_TOOL or status != "verified":
            skipped += 1
            continue
        users_by_email[email] = row
    return [users_by_email[email] for email in sorted(users_by_email)], skipped


def _contract_status(end_date: date) -> str:
    return "Expired" if end_date < date.today() else "Active"


async def import_autodesk_autocad(path: Path, *, apply: bool) -> AutodeskImportStats:
    users, skipped = _load_autocad_users(path)
    stats = AutodeskImportStats(autocad_users=len(users), skipped_rows=skipped)
    now = datetime.utcnow()

    async with SessionLocal() as session:
        existing_assignments = (await session.execute(select(LicenseAssignment))).scalars().all()
        for row in existing_assignments:
            if AUTOCAD_TOOL in {short_name_for(row.tool_name), short_name_for(row.tool_short_name)}:
                await session.delete(row)
                stats.deleted_assignments += 1

        existing_contracts = (await session.execute(select(LicenseContract))).scalars().all()
        for row in existing_contracts:
            if AUTOCAD_TOOL in {short_name_for(row.software), short_name_for(row.short_name)}:
                await session.delete(row)
                stats.deleted_contracts += 1

        await session.flush()

        for row in users:
            session.add(
                LicenseAssignment(
                    work_email=_email(row.get("email")),
                    tool_name=AUTOCAD_TOOL,
                    tool_short_name=AUTOCAD_TOOL,
                    plan="Autodesk single-user",
                    status="Assigned",
                    cost_centre=_text(row.get("team_alias")) or None,
                    notes=f"Synced from {path.name}; Autodesk ID {_text(row.get('autodesk_id'))}",
                    updated_at=now,
                )
            )
            stats.created_assignments += 1

        for contract_no, seats, end_date in AUTOCAD_CONTRACTS:
            status = _contract_status(end_date)
            if status == "Expired":
                stats.expired_contract_seats += seats
            else:
                stats.active_contract_seats += seats
            session.add(
                LicenseContract(
                    contract_key=f"ADSK-{contract_no}",
                    entity="IT -ADMIN",
                    software=AUTOCAD_TOOL,
                    short_name=AUTOCAD_TOOL,
                    category="Design Software",
                    contract_no=contract_no,
                    contract_type="1 year",
                    seats=seats,
                    vendor="Autodesk",
                    end_date=end_date,
                    status=status,
                    user_type="Single User",
                    notes=f"Synced from Autodesk subscriptions/contracts on {date.today().isoformat()}",
                    updated_at=now,
                )
            )
            stats.created_contracts += 1

        if apply:
            await session.commit()
        else:
            await session.rollback()

    return stats


async def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--file", required=True)
    parser.add_argument("--apply", action="store_true")
    args = parser.parse_args()

    path = Path(args.file).resolve()
    if not path.exists():
        raise FileNotFoundError(path)

    stats = await import_autodesk_autocad(path, apply=args.apply)
    print(("APPLIED" if args.apply else "DRY RUN") + f" AutoCAD import from {path}")
    print(f"  AutoCAD users: {stats.autocad_users}")
    print(f"  skipped non-AutoCAD/unverified rows: {stats.skipped_rows}")
    print(f"  deleted old AutoCAD assignments: {stats.deleted_assignments}")
    print(f"  created AutoCAD assignments: {stats.created_assignments}")
    print(f"  deleted old AutoCAD contracts: {stats.deleted_contracts}")
    print(f"  created AutoCAD contracts: {stats.created_contracts}")
    print(f"  active contract seats: {stats.active_contract_seats}")
    print(f"  expired contract seats: {stats.expired_contract_seats}")


if __name__ == "__main__":
    async def _run() -> None:
        try:
            await main()
        finally:
            await engine.dispose()

    asyncio.run(_run())
