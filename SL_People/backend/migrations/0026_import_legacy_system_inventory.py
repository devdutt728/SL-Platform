"""Import the legacy System Inventory workbook into current asset tables.

This handles the non-normalized workbook with sheets such as
`SYSTEM INVENTORY LIST`, `ALL SYSTEM SOFTWARE LIST`, and `LAPTOP`.
It intentionally ignores password/license-key columns and assigns installed
software to the current system holder email, not to shared account IDs embedded
in the workbook.

Usage:
    python -m migrations.0026_import_legacy_system_inventory --file "C:/Users/Dev/Downloads/Copy of System Inventory.xlsx"
    python -m migrations.0026_import_legacy_system_inventory --file "..." --apply
"""

from __future__ import annotations

import argparse
import asyncio
import csv
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any, Optional

from openpyxl import load_workbook
from sqlalchemy import func, select

from app.db.platform_session import PlatformSessionLocal, platform_engine
from app.db.session import SessionLocal, engine
from app.models.people import LicenseAssignment, OrgEmployee, OrgGroup, SystemInventory
from app.models.platform_person import DimPerson
from app.services.console_logic import grade_pc, short_name_for


SYSTEM_SHEET = "SYSTEM INVENTORY LIST"
SOFTWARE_SHEET = " ALL SYSTEM SOFTWARE LIST"
LAPTOP_SHEET = "LAPTOP"

SYSTEM_ALIASES = {
    "LDS-101/LUM": "LDS-15",
    "ACOOUNT ROOM": "ACOOUNTROOM",
    "ACCOUNT ROOM": "ACOOUNTROOM",
}

NO_VALUES = {"", "no", "none", "0", "n/a", "na", "nil", "-"}


@dataclass
class LegacySystem:
    system_id: str
    system_type: str
    source_row: str
    user_display: Optional[str] = None
    team: Optional[str] = None
    processor: Optional[str] = None
    ram_gb: Optional[float] = None
    ram_slots_free: Optional[str] = None
    graphics_card: Optional[str] = None
    cpu_cores: Optional[str] = None
    storage: Optional[str] = None
    motherboard: Optional[str] = None
    purchase_date: Optional[date] = None
    vendor: Optional[str] = None
    serial_no: Optional[str] = None
    service_tag: Optional[str] = None
    os: Optional[str] = None
    office_version: Optional[str] = None
    autocad_version: Optional[str] = None
    adobe_versions: Optional[str] = None
    sketchup_version: Optional[str] = None
    threedmax_version: Optional[str] = None
    rhino_version: Optional[str] = None
    enscape_version: Optional[str] = None
    d5_render: Optional[str] = None
    antivirus: Optional[str] = None


@dataclass
class ImportStats:
    systems_seen: int = 0
    systems_updated: int = 0
    systems_created: int = 0
    license_rows_added: int = 0
    license_rows_existing: int = 0
    skipped_unassigned_software: int = 0
    warnings: list[str] = field(default_factory=list)


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _none(value: Any) -> Optional[str]:
    text = _text(value)
    if not text or text.lower() in {"nan", "none"}:
        return None
    return text


def _date(value: Any) -> Optional[date]:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return None


def _canonical_key(value: Any) -> str:
    text = _text(value).upper().replace("\\", "/")
    text = re.sub(r"\s+", " ", text).strip()
    text = text.replace("LDS ", "LDS-")
    if re.match(r"^LAP\d+$", text):
        text = text.replace("LAP", "LAP ")
    text = SYSTEM_ALIASES.get(text, text)
    return re.sub(r"[^A-Z0-9/-]+", "", text)


def _is_present(value: Optional[str]) -> bool:
    return str(value or "").strip().lower() not in NO_VALUES


def _ram_gb(*values: Any) -> Optional[float]:
    for value in values:
        text = _text(value).upper()
        match = re.search(r"(\d+(?:\.\d+)?)\s*GB", text)
        if match:
            return float(match.group(1))
        match = re.search(r"\b(\d{1,3})\b", text)
        if match:
            return float(match.group(1))
    return None


def _join_versions(*values: Any) -> Optional[str]:
    parts: list[str] = []
    for value in values:
        text = _none(value)
        if text and _is_present(text) and text.lower() not in {p.lower() for p in parts}:
            parts.append(text)
    return " / ".join(parts) if parts else None


def _installed_tools(row: LegacySystem) -> list[str]:
    tools: list[str] = []
    if _is_present(row.office_version):
        office = row.office_version or ""
        tools.append("MS Office 365 Business Standard" if "business" in office.lower() else "MS Office 365 Apps")
    if _is_present(row.autocad_version):
        tools.append("AutoCAD LT")
    adobe = str(row.adobe_versions or "").lower()
    if "photoshop" in adobe:
        tools.append("Adobe Photoshop")
    if "illustrator" in adobe:
        tools.append("Adobe Illustrator")
    if "indesign" in adobe or "in design" in adobe:
        tools.append("Adobe InDesign")
    if "acrobat" in adobe:
        tools.append("Adobe Acrobat")
    if _is_present(row.sketchup_version):
        tools.append("SketchUp Pro")
    if _is_present(row.threedmax_version):
        tools.append("3ds Max")
    if _is_present(row.rhino_version):
        tools.append("Rhino")
    if _is_present(row.enscape_version):
        tools.append("Enscape")
    if _is_present(row.d5_render):
        tools.append("D5 Render")
    return list(dict.fromkeys(tools))


def _load_legacy_systems(path: Path) -> dict[str, LegacySystem]:
    wb = load_workbook(path, data_only=True)
    systems: dict[str, LegacySystem] = {}

    ws = wb[SYSTEM_SHEET]
    for row_idx in range(3, ws.max_row + 1):
        raw_id = _text(ws.cell(row_idx, 2).value)
        key = _canonical_key(raw_id)
        if not key or "👉" in raw_id:
            continue
        systems[key] = LegacySystem(
            system_id=raw_id,
            system_type="Desktop",
            source_row=f"{SYSTEM_SHEET}!{row_idx}",
            user_display=_none(ws.cell(row_idx, 3).value),
            team=_none(ws.cell(row_idx, 4).value),
            processor=_none(ws.cell(row_idx, 5).value),
            motherboard=_none(ws.cell(row_idx, 6).value),
            ram_gb=_ram_gb(ws.cell(row_idx, 8).value, ws.cell(row_idx, 7).value),
            ram_slots_free=_none(ws.cell(row_idx, 9).value),
            storage=_none(ws.cell(row_idx, 10).value),
            cpu_cores=_none(ws.cell(row_idx, 11).value),
            graphics_card=_none(ws.cell(row_idx, 12).value),
            purchase_date=_date(ws.cell(row_idx, 13).value),
            vendor=_none(ws.cell(row_idx, 14).value),
        )

    ws = wb[LAPTOP_SHEET]
    for row_idx in range(3, ws.max_row + 1):
        raw_id = _text(ws.cell(row_idx, 2).value) or _text(ws.cell(row_idx, 8).value)
        key = _canonical_key(raw_id)
        if not key:
            continue
        systems[key] = LegacySystem(
            system_id=raw_id,
            system_type="Laptop",
            source_row=f"{LAPTOP_SHEET}!{row_idx}",
            user_display=_none(ws.cell(row_idx, 3).value),
            team=_none(ws.cell(row_idx, 4).value),
            motherboard=_none(ws.cell(row_idx, 5).value),
            processor=_none(ws.cell(row_idx, 6).value),
            ram_gb=_ram_gb(ws.cell(row_idx, 7).value),
            serial_no=_none(ws.cell(row_idx, 8).value),
            service_tag=_none(ws.cell(row_idx, 9).value),
        )

    ws = wb[SOFTWARE_SHEET]
    for row_idx in range(2, ws.max_row + 1):
        key = _canonical_key(ws.cell(row_idx, 2).value)
        if not key:
            continue
        row = systems.get(key)
        if row is None:
            continue
        row.os = _none(ws.cell(row_idx, 5).value)
        row.office_version = _none(ws.cell(row_idx, 7).value)
        row.autocad_version = _none(ws.cell(row_idx, 8).value)
        row.adobe_versions = _join_versions(ws.cell(row_idx, 9).value, ws.cell(row_idx, 15).value)
        row.sketchup_version = _none(ws.cell(row_idx, 10).value)
        row.antivirus = _none(ws.cell(row_idx, 18).value)
        row.threedmax_version = _none(ws.cell(row_idx, 22).value)
        row.rhino_version = _none(ws.cell(row_idx, 28).value)
        row.enscape_version = _none(ws.cell(row_idx, 29).value)
        row.d5_render = _none(ws.cell(row_idx, 30).value)

    return systems


def _active_assignment(status: Optional[str]) -> bool:
    return str(status or "").strip().lower() not in {"revoked", "available", "unassigned", "inactive", "disabled", "removed"}


async def _apply_import(path: Path, *, apply: bool, report: Optional[Path]) -> ImportStats:
    legacy = _load_legacy_systems(path)
    stats = ImportStats(systems_seen=len(legacy))
    now = datetime.utcnow()

    async with SessionLocal() as session, PlatformSessionLocal() as platform_session:
        systems = (await session.execute(select(SystemInventory))).scalars().all()
        systems_by_key = {_canonical_key(row.system_id): row for row in systems}

        people = (await platform_session.execute(select(DimPerson))).scalars().all()
        people_by_email = {str(row.email or "").strip().lower(): row for row in people if row.email}
        people_by_code = {str(row.person_code or ""): row for row in people if row.person_code}

        org_rows = (await session.execute(select(OrgEmployee))).scalars().all()
        groups = (await session.execute(select(OrgGroup))).scalars().all()
        group_by_key = {row.group_key: row for row in groups}
        org_by_emp = {row.employee_no: row for row in org_rows if row.employee_no}

        license_rows = (await session.execute(select(LicenseAssignment))).scalars().all()
        active_license_keys = {
            (str(row.work_email or "").strip().lower(), short_name_for(row.tool_short_name or row.tool_name))
            for row in license_rows
            if _active_assignment(row.status)
        }

        report_rows: list[dict[str, Any]] = []
        for key, legacy_row in legacy.items():
            system = systems_by_key.get(key)
            created = False
            if system is None:
                system = SystemInventory(system_id=SYSTEM_ALIASES.get(key, legacy_row.system_id) or key)
                created = True
                systems_by_key[key] = system
                session.add(system)

            current_email = str(system.assigned_email or "").strip().lower()
            person = people_by_email.get(current_email)
            org = org_by_emp.get(str(person.person_code)) if person and person.person_code else None
            group = group_by_key.get(org.group_key) if org else None

            system.system_type = legacy_row.system_type or system.system_type
            system.user_display = (person.full_name or person.display_name) if person else (legacy_row.user_display or system.user_display)
            system.team = group.name if group else (legacy_row.team or system.team)
            system.processor = legacy_row.processor or system.processor
            system.ram_gb = legacy_row.ram_gb if legacy_row.ram_gb is not None else system.ram_gb
            system.ram_slots_free = legacy_row.ram_slots_free or system.ram_slots_free
            system.graphics_card = legacy_row.graphics_card or system.graphics_card
            system.cpu_cores = legacy_row.cpu_cores or system.cpu_cores
            system.storage = legacy_row.storage or system.storage
            system.motherboard = legacy_row.motherboard or system.motherboard
            system.os = legacy_row.os or system.os
            system.office_version = legacy_row.office_version or system.office_version
            system.autocad_version = legacy_row.autocad_version or system.autocad_version
            system.adobe_versions = legacy_row.adobe_versions or system.adobe_versions
            system.sketchup_version = legacy_row.sketchup_version or system.sketchup_version
            system.threedmax_version = legacy_row.threedmax_version or system.threedmax_version
            system.rhino_version = legacy_row.rhino_version or system.rhino_version
            system.enscape_version = legacy_row.enscape_version or system.enscape_version
            system.d5_render = legacy_row.d5_render or system.d5_render
            system.antivirus = legacy_row.antivirus or system.antivirus
            system.purchase_date = legacy_row.purchase_date or system.purchase_date
            system.vendor = legacy_row.vendor or system.vendor
            system.service_tag = legacy_row.service_tag or system.service_tag
            system.serial_no = legacy_row.serial_no or system.serial_no
            system.status = system.status or "Active"

            grade = grade_pc(system.processor, system.graphics_card, str(system.ram_gb or ""))
            system.composite_score = grade.score
            system.capability_tier = grade.tier
            system.upgrade_suggestion = grade.suggestion
            system.updated_at = now

            if created:
                stats.systems_created += 1
            else:
                stats.systems_updated += 1

            installed = _installed_tools(legacy_row)
            if installed and not current_email:
                stats.skipped_unassigned_software += len(installed)
                stats.warnings.append(f"{system.system_id}: software present but no assigned_email; skipped license holder updates")
            for tool in installed:
                if not current_email:
                    continue
                license_key = (current_email, tool)
                if license_key in active_license_keys:
                    stats.license_rows_existing += 1
                    continue
                assignment = LicenseAssignment(
                    work_email=current_email,
                    tool_name=tool,
                    tool_short_name=short_name_for(tool),
                    plan="Imported from installed software",
                    status="Assigned",
                    cost_centre=system.team,
                    notes=f"Derived from {path.name} {legacy_row.source_row} for {system.system_id}",
                    updated_at=now,
                )
                session.add(assignment)
                active_license_keys.add(license_key)
                stats.license_rows_added += 1

            report_rows.append(
                {
                    "system_id": system.system_id,
                    "source_row": legacy_row.source_row,
                    "assigned_email": current_email,
                    "user_display": system.user_display,
                    "team": system.team,
                    "score": system.composite_score,
                    "tier": system.capability_tier,
                    "installed_tools": "; ".join(installed),
                    "created": created,
                }
            )

        if report:
            report.parent.mkdir(parents=True, exist_ok=True)
            with report.open("w", newline="", encoding="utf-8") as fh:
                writer = csv.DictWriter(fh, fieldnames=list(report_rows[0].keys()) if report_rows else [])
                if report_rows:
                    writer.writeheader()
                    writer.writerows(report_rows)

        if apply:
            await session.commit()
        else:
            await session.rollback()

    return stats


async def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--file", required=True)
    parser.add_argument("--apply", action="store_true")
    parser.add_argument("--report", default="")
    args = parser.parse_args()

    path = Path(args.file).resolve()
    if not path.exists():
        raise FileNotFoundError(path)
    report = Path(args.report).resolve() if args.report else None
    stats = await _apply_import(path, apply=args.apply, report=report)
    mode = "APPLIED" if args.apply else "DRY RUN"
    print(f"{mode} legacy inventory import from {path}")
    print(f"  systems seen: {stats.systems_seen}")
    print(f"  systems updated: {stats.systems_updated}")
    print(f"  systems created: {stats.systems_created}")
    print(f"  license rows added: {stats.license_rows_added}")
    print(f"  license rows already present: {stats.license_rows_existing}")
    print(f"  skipped unassigned software licenses: {stats.skipped_unassigned_software}")
    if report:
        print(f"  report: {report}")
    if stats.warnings:
        print("  warnings:")
        for warning in stats.warnings[:30]:
            print(f"    - {warning}")


if __name__ == "__main__":
    async def _run() -> None:
        try:
            await main()
        finally:
            await engine.dispose()
            await platform_engine.dispose()

    asyncio.run(_run())
