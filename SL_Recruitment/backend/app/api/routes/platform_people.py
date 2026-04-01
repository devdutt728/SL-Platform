from __future__ import annotations

import csv
import hashlib
import io
import json
import re
from dataclasses import dataclass
from datetime import date, datetime
from typing import Any

from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, UploadFile, status
from sqlalchemy import Integer, delete, func, or_, select, text
from sqlalchemy.exc import SQLAlchemyError
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.auth import require_roles, require_superadmin
from app.core.roles import Role
from app.db.platform_session import get_platform_session
from app.models.platform_person import DimPerson, DimPersonExtra, DimPersonRole
from app.models.platform_role import DimRole
from app.schemas.platform import (
    BulkUploadError,
    BulkUploadResult,
    BulkUploadWarning,
    PlatformPersonCreate,
    PlatformPersonOut,
    PlatformPersonSuggestion,
    PlatformPersonUpdate,
)
from app.schemas.user import UserContext
from app.services.platform_identity import active_status_filter

router = APIRouter(prefix="/platform/people", tags=["platform"])

_PERSON_CORE_FIELDS = {
    "person_id",
    "person_code",
    "personal_id",
    "first_name",
    "last_name",
    "email",
    "mobile_number",
    "role_id",
    "grade_id",
    "department_id",
    "manager_id",
    "employment_type",
    "join_date",
    "exit_date",
    "status",
    "is_deleted",
    "created_at",
    "updated_at",
    "source_system",
    "source_candidate_id",
    "source_candidate_code",
    "full_name",
    "display_name",
    "middle_name",
    "date_of_birth",
    "gender",
    "marital_status",
    "marriage_date",
    "blood_group",
    "physically_handicapped",
    "nationality",
    "work_phone",
    "home_phone",
    "personal_email",
    "current_address_line_1",
    "current_address_line_2",
    "current_address_city",
    "current_address_state",
    "current_address_zip",
    "current_address_country",
    "permanent_address_line_1",
    "permanent_address_line_2",
    "permanent_address_city",
    "permanent_address_state",
    "permanent_address_zip",
    "permanent_address_country",
    "father_name",
    "mother_name",
    "spouse_name",
    "children_names",
    "attendance_number",
    "location",
    "location_country",
    "legal_entity",
    "business_unit",
    "department",
    "sub_department",
    "job_title",
    "secondary_job_title",
    "reporting_to",
    "dotted_line_manager",
    "leave_plan",
    "band",
    "pay_grade",
    "time_type",
    "shift_policy_name",
    "weekly_off_policy_name",
    "attendance_time_tracking_policy",
    "attendance_capture_scheme",
    "holiday_list_name",
    "expense_policy_name",
    "notice_period",
    "aadhaar_number",
    "pf_number",
    "uan_number",
    "comments",
    "exit_status",
    "termination_type",
    "termination_reason",
    "resignation_note",
    "cost_center",
}

_HEADER_ALIASES = {
    "employee_number": "person_code",
    "work_email": "email",
    "mobile_phone": "mobile_number",
    "reporting_manager_employee_number": "manager_id",
    "date_joined": "join_date",
    "employment_status": "status",
    "worker_type": "employment_type",
    "pan_card_number": "personal_id",
    "pan_number": "personal_id",
}

_MANDATORY_FIELDS = {"person_code", "first_name", "email"}
_UPDATABLE_TO_NONE_FIELDS = {"exit_date", "is_deleted"}
_NULLISH_TEXT = {"na", "n/a", "none", "null", "not_available", "not_applicable", "-"}
_REPLACE_CONFIRM_TOKEN = "REPLACE_DIM_PERSON"
_DEFAULT_EXTERNAL_IDENTITY_SOURCE = "emp_master_upload"


@dataclass
class _BulkRow:
    row_number: int
    normalized: dict[str, Any]


@dataclass
class _PendingManagerAssignment:
    row_number: int
    person: DimPerson
    normalized: dict[str, Any]
    was_unchanged: bool


@dataclass
class _PersonCodeCounterState:
    code_type: str
    next_seq: int
    original_next_seq: int
    max_seq: int | None = None


def _external_identity_source_from_payload(payload: dict[str, Any]) -> str:
    source = _normalize_text(payload.get("external_identity_source")) or _normalize_text(payload.get("source_system"))
    return (source or _DEFAULT_EXTERNAL_IDENTITY_SOURCE).strip().lower()


def _external_identity_key(source: str, external_employee_code: str) -> str:
    return f"{source.strip().lower()}::{external_employee_code.strip().lower()}"


def _external_employee_code(normalized: dict[str, Any]) -> str:
    return _normalize_text(normalized.get("external_employee_code"))


def _external_manager_code(normalized: dict[str, Any]) -> str:
    return _normalize_text(normalized.get("external_manager_code"))


def _incoming_internal_person_code(normalized: dict[str, Any]) -> str:
    person_code = _normalize_text(normalized.get("person_code"))
    external_code = _external_employee_code(normalized)
    if person_code and external_code and person_code == external_code:
        return ""
    return person_code


def _missing_bulk_required_fields(normalized: dict[str, Any]) -> list[str]:
    missing: list[str] = []
    if not (_incoming_internal_person_code(normalized) or _external_employee_code(normalized)):
        missing.append("person_code")
    if not _normalize_text(normalized.get("first_name")):
        missing.append("first_name")
    if not _normalize_email(normalized.get("email")):
        missing.append("email")
    return missing

def _is_superadmin_user(user: UserContext) -> bool:
    if Role.HR_ADMIN in (getattr(user, "roles", None) or []):
        return True
    if (getattr(user, "platform_role_id", None) or None) == 2:
        return True
    if (getattr(user, "platform_role_code", None) or "").strip() == "2":
        return True
    role_ids = getattr(user, "platform_role_ids", None) or []
    if 2 in role_ids:
        return True
    role_codes = getattr(user, "platform_role_codes", None) or []
    normalized_codes = {str(code).strip().lower() for code in role_codes if code}
    return "s_admin" in normalized_codes or "superadmin" in normalized_codes


@router.get("", response_model=list[PlatformPersonSuggestion])
async def search_people(
    q: str = Query(default="", max_length=100),
    limit: int = Query(default=10, ge=1, le=25),
    include_deleted: bool = Query(default=False),
    session: AsyncSession = Depends(get_platform_session),
    _user: UserContext = Depends(require_roles([Role.HR_ADMIN, Role.HR_EXEC, Role.HIRING_MANAGER])),
):
    q_norm = q.strip().lower()
    like = f"%{q_norm}%"

    is_superadmin = _is_superadmin_user(_user)
    base_filters = [(DimPerson.is_deleted == 0) | (DimPerson.is_deleted.is_(None))]
    if is_superadmin and include_deleted:
        base_filters = []
    if not is_superadmin:
        base_filters.append(active_status_filter())

    # If query is too short, return a small "default" list for dropdown convenience.
    if len(q_norm) < 2:
        rows = (
            await session.execute(
                select(
                    DimPerson.person_id,
                    DimPerson.person_code,
                    DimPerson.email,
                    DimPerson.first_name,
                    DimPerson.last_name,
                    DimPerson.display_name,
                    DimPerson.full_name,
                    DimPerson.status,
                    DimPerson.is_deleted,
                    DimPerson.role_id,
                    DimRole.role_code,
                    DimRole.role_name,
                )
                .select_from(DimPerson)
                .outerjoin(DimRole, DimRole.role_id == DimPerson.role_id)
                .where(*base_filters)
                .order_by(func.coalesce(DimPerson.display_name, DimPerson.full_name, DimPerson.first_name, DimPerson.email).asc())
                .limit(limit)
            )
        ).all()
    else:
        rows = (
            await session.execute(
                select(
                    DimPerson.person_id,
                    DimPerson.person_code,
                    DimPerson.email,
                    DimPerson.first_name,
                    DimPerson.last_name,
                    DimPerson.display_name,
                    DimPerson.full_name,
                    DimPerson.status,
                    DimPerson.is_deleted,
                    DimPerson.role_id,
                    DimRole.role_code,
                    DimRole.role_name,
                )
                .select_from(DimPerson)
                .outerjoin(DimRole, DimRole.role_id == DimPerson.role_id)
                .where(
                    *base_filters,
                    or_(
                        func.lower(DimPerson.email).like(like),
                        func.lower(func.coalesce(DimPerson.display_name, "")).like(like),
                        func.lower(func.coalesce(DimPerson.full_name, "")).like(like),
                        func.lower(func.coalesce(DimPerson.first_name, "")).like(like),
                        func.lower(func.coalesce(DimPerson.last_name, "")).like(like),
                        func.lower(func.coalesce(DimPerson.person_id, "")).like(like),
                        func.lower(func.coalesce(DimPerson.person_code, "")).like(like),
                        func.lower(func.coalesce(DimPerson.personal_id, "")).like(like),
                        func.lower(
                            func.concat(
                                func.coalesce(DimPerson.first_name, ""),
                                " ",
                                func.coalesce(DimPerson.last_name, ""),
                            )
                        ).like(like),
                    ),
                )
                .limit(limit)
            )
        ).all()

    person_ids = [row.person_id for row in rows]
    roles_map: dict[str, dict[str, list]] = {}
    if person_ids:
        role_rows = (
            await session.execute(
                select(
                    DimPersonRole.person_id,
                    DimRole.role_id,
                    DimRole.role_code,
                    DimRole.role_name,
                )
                .select_from(DimPersonRole)
                .join(DimRole, DimRole.role_id == DimPersonRole.role_id)
                .where(DimPersonRole.person_id.in_(person_ids))
                .order_by(DimPersonRole.person_id.asc(), DimRole.role_id.asc())
            )
        ).all()
        for role_row in role_rows:
            bucket = roles_map.setdefault(
                role_row.person_id, {"role_ids": [], "role_codes": [], "role_names": []}
            )
            if role_row.role_id is not None:
                bucket["role_ids"].append(int(role_row.role_id))
            if role_row.role_code:
                bucket["role_codes"].append(str(role_row.role_code))
            if role_row.role_name:
                bucket["role_names"].append(str(role_row.role_name))

    out: list[PlatformPersonSuggestion] = []
    for row in rows:
        first = (row.first_name or "").strip()
        last = (row.last_name or "").strip()
        full_name = (row.display_name or row.full_name or f"{first} {last}").strip() or row.email
        role_bucket = roles_map.get(row.person_id, {"role_ids": [], "role_codes": [], "role_names": []})
        if not role_bucket["role_ids"] and row.role_id is not None:
            role_bucket = {
                "role_ids": [int(row.role_id)],
                "role_codes": [str(row.role_code)] if row.role_code else [],
                "role_names": [str(row.role_name)] if row.role_name else [],
            }
        out.append(
            PlatformPersonSuggestion(
                person_id=row.person_id,
                person_code=row.person_code,
                full_name=full_name,
                email=row.email,
                status=row.status,
                is_deleted=row.is_deleted,
                role_code=row.role_code,
                role_name=row.role_name,
                role_ids=role_bucket["role_ids"],
                role_codes=role_bucket["role_codes"],
                role_names=role_bucket["role_names"],
            )
        )
    return out


@router.post("/bulk", response_model=BulkUploadResult)
async def bulk_upload_people(
    file: UploadFile = File(...),
    overwrite: bool = Form(default=False),
    replace_all: bool = Form(default=False),
    confirm_replace: str | None = Form(default=None),
    dry_run: bool = Form(default=False),
    session: AsyncSession = Depends(get_platform_session),
    _user: UserContext = Depends(require_superadmin()),
):
    if not file.filename:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Missing file")
    if replace_all and not dry_run and (confirm_replace or "").strip() != _REPLACE_CONFIRM_TOKEN:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"replace_all requires confirm_replace={_REPLACE_CONFIRM_TOKEN}",
        )

    raw = await file.read()
    if not raw:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="File is empty")

    if overwrite:
        # Legacy flag retained for backward compatibility. Safe merge is always enabled now.
        overwrite = False

    rows = _load_bulk_rows(file.filename, raw)
    rows = _dedupe_bulk_rows(rows)
    result = BulkUploadResult(
        mode="dry_run" if dry_run else "apply",
        batch_hash=hashlib.sha256(raw).hexdigest(),
        total=len(rows),
        created=0,
        updated=0,
        skipped=0,
        unchanged=0,
        conflicts=0,
        errors=[],
        warnings=[],
    )
    if not rows:
        return result

    if not dry_run:
        await _ensure_person_extra_table(session)

    if replace_all:
        return await _replace_all_people(
            session=session,
            rows=rows,
            result=result,
            dry_run=dry_run,
        )

    existing_people = await _load_existing_people(session, rows)
    by_person_id, by_person_code, by_email = _build_people_maps(existing_people)
    by_external_employee_code = await _load_external_employee_identity_map(
        session=session,
        by_person_id=by_person_id,
        warnings=result.warnings,
    )
    person_code_counters = await _lock_person_code_counters(session)
    next_seq: int | None = None
    used_ids = {person.person_id for person in existing_people if (person.person_id or "").strip()}
    used_ids.discard("")
    preserved_person_code_count = 0
    unresolved_manager_count = 0
    pending_manager_assignments: list[_PendingManagerAssignment] = []

    for row in rows:
        normalized = row.normalized
        _apply_external_identity_defaults(normalized)
        row_number = row.row_number
        external_employee_code = _external_employee_code(normalized)
        person_code = _incoming_internal_person_code(normalized) or external_employee_code
        first_name = _normalize_text(normalized.get("first_name"))
        email = _normalize_email(normalized.get("email"))
        incoming_person_id = _normalize_text(normalized.get("person_id"))

        if _incoming_internal_person_code(normalized):
            normalized["person_code"] = person_code
        if first_name:
            normalized["first_name"] = first_name
        if email:
            normalized["email"] = email
        if incoming_person_id:
            normalized["person_id"] = incoming_person_id

        missing = _missing_bulk_required_fields(normalized)
        if missing:
            result.skipped += 1
            result.errors.append(
                BulkUploadError(
                    row=row_number,
                    message=f"Missing required field(s): {', '.join(sorted(missing))}",
                    person_id=incoming_person_id,
                    person_code=person_code,
                    email=email,
                )
            )
            continue

        try:
            existing = _resolve_existing_person(
                row=row,
                by_person_id=by_person_id,
                by_person_code=by_person_code,
                by_external_employee_code=by_external_employee_code,
                by_email=by_email,
                warnings=result.warnings,
            )
            updates = _apply_status_defaults(_apply_name_defaults(_apply_row_updates(normalized)))

            if existing is not None:
                if _preserve_existing_person_code(
                    existing=existing,
                    updates=updates,
                ):
                    preserved_person_code_count += 1
                if incoming_person_id and incoming_person_id != existing.person_id:
                    result.warnings.append(
                        BulkUploadWarning(
                            row=row_number,
                            message=(
                                "Incoming person_id does not match resolved record. Existing person_id kept."
                            ),
                            person_id=existing.person_id,
                            person_code=person_code,
                            email=email,
                        )
                    )
                updates.pop("person_id", None)
                patch = _build_safe_update_payload(existing, updates)
                patch.pop("manager_id", None)
                _assert_no_update_conflicts(existing, patch, by_person_code, by_email)
                if not patch:
                    _refresh_external_employee_identity_map(
                        person=existing,
                        normalized=normalized,
                        by_external_employee_code=by_external_employee_code,
                    )
                    pending_manager_assignments.append(
                        _PendingManagerAssignment(
                            row_number=row_number,
                            person=existing,
                            normalized=dict(normalized),
                            was_unchanged=True,
                        )
                    )
                    await _upsert_person_extra(
                        session=session,
                        person=existing,
                        normalized=normalized,
                        batch_hash=result.batch_hash,
                        dry_run=dry_run,
                    )
                    result.unchanged += 1
                    continue
                if not dry_run:
                    for key, value in patch.items():
                        setattr(existing, key, value)
                    if "updated_at" not in patch:
                        existing.updated_at = datetime.utcnow()
                _refresh_person_maps(
                    person=existing,
                    by_person_id=by_person_id,
                    by_person_code=by_person_code,
                    by_email=by_email,
                )
                _refresh_external_employee_identity_map(
                    person=existing,
                    normalized=normalized,
                    by_external_employee_code=by_external_employee_code,
                )
                pending_manager_assignments.append(
                    _PendingManagerAssignment(
                        row_number=row_number,
                        person=existing,
                        normalized=dict(normalized),
                        was_unchanged=False,
                    )
                )
                await _upsert_person_extra(
                    session=session,
                    person=existing,
                    normalized=normalized,
                    batch_hash=result.batch_hash,
                    dry_run=dry_run,
                )
                result.updated += 1
                continue

            create_payload = updates
            create_payload.pop("person_id", None)
            create_payload["manager_id"] = None

            resolved_person_code, normalized_employment_type = await _resolve_bulk_person_code(
                session=session,
                normalized=normalized,
                payload=create_payload,
                by_person_code=by_person_code,
                counters=person_code_counters,
            )
            create_payload["person_code"] = resolved_person_code
            create_payload["employment_type"] = normalized_employment_type

            if not create_payload.get("source_system"):
                create_payload["source_system"] = _DEFAULT_EXTERNAL_IDENTITY_SOURCE
            if not create_payload.get("person_id"):
                next_seq = await _ensure_next_seq(session, next_seq)
                seq_value = _consume_next_seq(_build_person_prefix(first_name, _normalize_text(normalized.get("last_name"))), next_seq, used_ids)
                create_payload["person_id"] = _format_person_id(first_name, _normalize_text(normalized.get("last_name")), seq_value)
                next_seq += 1
            person_id = _normalize_text(create_payload.get("person_id"))
            if not person_id:
                raise ValueError("Unable to generate person_id")
            if person_id in by_person_id:
                raise ValueError(f"person_id '{person_id}' already exists")
            create_payload["person_id"] = person_id

            _assert_no_create_conflicts(create_payload, by_person_code, by_email)
            person = DimPerson(**create_payload)
            if person.created_at is None:
                person.created_at = datetime.utcnow()
            if person.updated_at is None:
                person.updated_at = datetime.utcnow()
            if not dry_run:
                session.add(person)
            _refresh_person_maps(
                person=person,
                by_person_id=by_person_id,
                by_person_code=by_person_code,
                by_email=by_email,
            )
            _refresh_external_employee_identity_map(
                person=person,
                normalized=normalized,
                by_external_employee_code=by_external_employee_code,
            )
            pending_manager_assignments.append(
                _PendingManagerAssignment(
                    row_number=row_number,
                    person=person,
                    normalized=dict(normalized),
                    was_unchanged=False,
                )
            )
            await _upsert_person_extra(
                session=session,
                person=person,
                normalized=normalized,
                batch_hash=result.batch_hash,
                dry_run=dry_run,
            )
            used_ids.add(person.person_id)
            result.created += 1
        except Exception as exc:
            result.conflicts += 1
            result.skipped += 1
            result.errors.append(
                BulkUploadError(
                    row=row_number,
                    message=str(exc),
                    person_id=incoming_person_id,
                    person_code=person_code,
                    email=email,
                )
            )

    for pending in pending_manager_assignments:
        requested_manager_token = _requested_manager_reference(pending.normalized)
        if not requested_manager_token:
            continue
        manager_id = _resolve_manager_reference(
            normalized=pending.normalized,
            by_person_id=by_person_id,
            by_person_code=by_person_code,
            by_external_employee_code=by_external_employee_code,
        )
        if manager_id == pending.person.person_id:
            manager_id = None
        if manager_id is None:
            unresolved_manager_count += 1
        current_manager_id = _normalize_text(pending.person.manager_id)
        next_manager_id = _normalize_text(manager_id)
        if current_manager_id == next_manager_id:
            continue
        if not dry_run:
            pending.person.manager_id = manager_id
            pending.person.updated_at = datetime.utcnow()
        if pending.was_unchanged:
            result.unchanged -= 1
            result.updated += 1

    if preserved_person_code_count:
        result.warnings.append(
            BulkUploadWarning(
                message=(
                    f"Safe merge kept the existing person_code for {preserved_person_code_count} "
                    "matched record(s). Use manual edits if you need to rewrite person_code values."
                )
            )
        )
    if unresolved_manager_count:
        result.warnings.append(
            BulkUploadWarning(
                message=(
                    f"Bulk import cleared unresolved manager references for {unresolved_manager_count} "
                    "row(s) because manager_id must reference an existing person_id."
                )
            )
        )

    result.processed = result.created + result.updated + result.unchanged
    if dry_run:
        await session.rollback()
    else:
        await _save_person_code_counters(session, person_code_counters)
        await session.commit()
    return result


async def _replace_all_people(
    *,
    session: AsyncSession,
    rows: list[_BulkRow],
    result: BulkUploadResult,
    dry_run: bool,
) -> BulkUploadResult:
    existing_people = (await session.execute(select(DimPerson))).scalars().all()
    legacy_by_code, legacy_by_email = _build_legacy_person_maps(existing_people, result.warnings)
    used_ids = {person.person_id for person in existing_people if (person.person_id or "").strip()}
    used_ids.discard("")
    next_seq = await _ensure_next_seq(session, None)
    now = datetime.utcnow()

    staged_people: list[DimPerson] = []
    staged_extras: list[DimPersonExtra] = []
    staged_by_person_id: dict[str, DimPerson] = {}
    staged_by_person_code: dict[str, DimPerson] = {}
    staged_by_email: dict[str, DimPerson] = {}
    staged_manager_tokens: dict[str, str] = {}
    staged_row_by_person_id: dict[str, int] = {}

    for row in rows:
        normalized = row.normalized
        row_number = row.row_number
        person_code = _normalize_text(normalized.get("person_code"))
        first_name = _normalize_text(normalized.get("first_name"))
        last_name = _normalize_text(normalized.get("last_name"))
        email = _normalize_email(normalized.get("email"))
        incoming_person_id = _normalize_text(normalized.get("person_id"))

        if person_code:
            normalized["person_code"] = person_code
        if first_name:
            normalized["first_name"] = first_name
        if email:
            normalized["email"] = email
        if incoming_person_id:
            normalized["person_id"] = incoming_person_id

        missing = [field for field in _MANDATORY_FIELDS if not _normalize_text(normalized.get(field))]
        if missing:
            result.skipped += 1
            result.errors.append(
                BulkUploadError(
                    row=row_number,
                    message=f"Missing required field(s): {', '.join(sorted(missing))}",
                    person_id=incoming_person_id,
                    person_code=person_code,
                    email=email,
                )
            )
            continue

        try:
            create_payload = _apply_status_defaults(_apply_name_defaults(_apply_row_updates(normalized)))
            manager_token = _normalize_text(create_payload.get("manager_id"))
            create_payload.pop("person_id", None)
            create_payload["person_id"], next_seq = _resolve_replace_person_id(
                row=row,
                incoming_person_id=incoming_person_id,
                person_code=person_code,
                email=email,
                first_name=first_name,
                last_name=last_name,
                legacy_by_code=legacy_by_code,
                legacy_by_email=legacy_by_email,
                used_ids=used_ids,
                next_seq=next_seq,
                warnings=result.warnings,
            )
            create_payload["manager_id"] = None
            if manager_token:
                staged_manager_tokens[create_payload["person_id"]] = manager_token
            if not create_payload.get("source_system"):
                create_payload["source_system"] = "emp_master_upload"
            if not _normalize_text(create_payload.get("status")):
                create_payload["status"] = "working"
            if create_payload.get("is_deleted") is None:
                create_payload["is_deleted"] = _derive_is_deleted_from_status(create_payload.get("status"))
            if create_payload.get("created_at") is None:
                create_payload["created_at"] = now
            if create_payload.get("updated_at") is None:
                create_payload["updated_at"] = now

            _assert_no_replace_staging_conflicts(
                payload=create_payload,
                by_person_id=staged_by_person_id,
                by_person_code=staged_by_person_code,
                by_email=staged_by_email,
            )
            person = DimPerson(**create_payload)
            staged_people.append(person)
            extra_record = _build_person_extra_record(
                person=person,
                normalized=normalized,
                batch_hash=result.batch_hash,
                now=now,
            )
            if extra_record is not None:
                staged_extras.append(extra_record)
            staged_by_person_id[person.person_id] = person
            person_code_key = _normalize_text(person.person_code).lower()
            if person_code_key:
                staged_by_person_code[person_code_key] = person
            email_key = _normalize_email(person.email)
            if email_key:
                staged_by_email[email_key] = person
            staged_row_by_person_id[person.person_id] = row_number
        except Exception as exc:
            result.conflicts += 1
            result.skipped += 1
            result.errors.append(
                BulkUploadError(
                    row=row_number,
                    message=str(exc),
                    person_id=incoming_person_id,
                    person_code=person_code,
                    email=email,
                )
            )

    manager_by_person_id = {key: [value] for key, value in staged_by_person_id.items()}
    manager_by_person_code = {key: [value] for key, value in staged_by_person_code.items()}
    for person in staged_people:
        manager_token = staged_manager_tokens.get(person.person_id, "")
        manager_id = _resolve_manager_id(
            manager_token,
            manager_by_person_id,
            manager_by_person_code,
        )
        if manager_id and manager_id not in staged_by_person_id:
            manager_id = None
        if manager_id == person.person_id:
            manager_id = None
        if manager_token and not manager_id:
            result.warnings.append(
                BulkUploadWarning(
                    row=staged_row_by_person_id.get(person.person_id),
                    message="Manager reference not found in upload set. manager_id cleared.",
                    person_id=person.person_id,
                    person_code=person.person_code,
                    email=person.email,
                )
            )
        person.manager_id = manager_id

    result.created = len(staged_people)
    result.updated = 0
    result.unchanged = 0
    result.processed = result.created

    if dry_run:
        await session.rollback()
        return result

    if not staged_people:
        await session.rollback()
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="replace_all aborted because no valid rows were found in the upload.",
        )

    await session.execute(delete(DimPersonRole))
    await session.execute(delete(DimPersonExtra))
    await session.execute(delete(DimPerson))
    if staged_people:
        session.add_all(staged_people)
        if staged_extras:
            session.add_all(staged_extras)
        role_rows = _build_person_role_rows(staged_people)
        if role_rows:
            session.add_all(role_rows)
    await session.commit()
    return result


def _build_legacy_person_maps(
    people: list[DimPerson],
    warnings: list[BulkUploadWarning],
) -> tuple[dict[str, str], dict[str, str]]:
    code_buckets: dict[str, set[str]] = {}
    email_buckets: dict[str, set[str]] = {}
    for person in people:
        code = _normalize_text(person.person_code).lower()
        if code:
            code_buckets.setdefault(code, set()).add(person.person_id)
        email = _normalize_email(person.email)
        if email:
            email_buckets.setdefault(email, set()).add(person.person_id)

    legacy_by_code: dict[str, str] = {}
    for code, ids in code_buckets.items():
        if len(ids) == 1:
            legacy_by_code[code] = next(iter(ids))
        else:
            warnings.append(
                BulkUploadWarning(
                    message=(
                        f"Existing table has duplicate person_code '{code}'. "
                        "Stable person_id preservation for this code is ambiguous."
                    ),
                )
            )

    legacy_by_email: dict[str, str] = {}
    for email, ids in email_buckets.items():
        if len(ids) == 1:
            legacy_by_email[email] = next(iter(ids))
        else:
            warnings.append(
                BulkUploadWarning(
                    message=(
                        f"Existing table has duplicate email '{email}'. "
                        "Stable person_id preservation for this email is ambiguous."
                    ),
                )
            )
    return legacy_by_code, legacy_by_email


def _resolve_replace_person_id(
    *,
    row: _BulkRow,
    incoming_person_id: str,
    person_code: str,
    email: str,
    first_name: str,
    last_name: str,
    legacy_by_code: dict[str, str],
    legacy_by_email: dict[str, str],
    used_ids: set[str],
    next_seq: int,
    warnings: list[BulkUploadWarning],
) -> tuple[str, int]:
    code_key = person_code.lower()
    preserved_id = legacy_by_code.get(code_key) if code_key else None
    if not preserved_id and email:
        preserved_id = legacy_by_email.get(email)

    if preserved_id:
        if incoming_person_id and incoming_person_id != preserved_id:
            warnings.append(
                BulkUploadWarning(
                    row=row.row_number,
                    message=(
                        "Incoming person_id differs from existing mapping. "
                        "Existing person_id preserved for stability."
                    ),
                    person_id=preserved_id,
                    person_code=person_code,
                    email=email,
                )
            )
        used_ids.add(preserved_id)
        return preserved_id, next_seq

    if incoming_person_id and not incoming_person_id.isdigit() and incoming_person_id not in used_ids:
        used_ids.add(incoming_person_id)
        return incoming_person_id, next_seq

    stable_from_code = _stable_person_id_from_code(person_code, used_ids)
    if stable_from_code:
        return stable_from_code, next_seq

    seq_value = _consume_next_seq(_build_person_prefix(first_name, last_name), next_seq, used_ids)
    return _format_person_id(first_name, last_name, seq_value), seq_value + 1


def _stable_person_id_from_code(person_code: str, used_ids: set[str]) -> str | None:
    raw = _normalize_text(person_code)
    if not raw:
        return None
    normalized = re.sub(r"[^a-z0-9]+", "_", raw.strip().lower()).strip("_")
    if not normalized:
        return None
    base = f"PC_{normalized.upper()}"
    base = base[:64].rstrip("_")
    if base and base not in used_ids:
        used_ids.add(base)
        return base

    digest = hashlib.sha1(normalized.encode("utf-8")).hexdigest()[:8].upper()
    suffix = f"_{digest}"
    max_base_len = 64 - len(suffix)
    candidate = f"{base[:max_base_len].rstrip('_')}{suffix}"
    if candidate not in used_ids:
        used_ids.add(candidate)
        return candidate

    index = 2
    while True:
        extra = f"{suffix}{index}"
        max_len = 64 - len(extra)
        candidate = f"{base[:max_len].rstrip('_')}{extra}"
        if candidate not in used_ids:
            used_ids.add(candidate)
            return candidate
        index += 1


def _assert_no_replace_staging_conflicts(
    *,
    payload: dict[str, Any],
    by_person_id: dict[str, DimPerson],
    by_person_code: dict[str, DimPerson],
    by_email: dict[str, DimPerson],
) -> None:
    person_id = _normalize_text(payload.get("person_id"))
    person_code = _normalize_text(payload.get("person_code")).lower()
    email = _normalize_email(payload.get("email"))
    if person_id and person_id in by_person_id:
        raise ValueError(f"person_id '{person_id}' is duplicated in upload.")
    if person_code and person_code in by_person_code:
        raise ValueError(f"person_code '{person_code}' is duplicated in upload.")
    if email and email in by_email:
        raise ValueError(f"email '{email}' is duplicated in upload.")


def _derive_is_deleted_from_status(status_value: Any) -> int:
    status = _normalize_text(status_value).lower()
    inactive_states = {
        "relieved",
        "inactive",
        "resigned",
        "terminated",
        "separated",
        "left",
        "exit",
    }
    return 1 if status in inactive_states else 0


def _build_person_role_rows(people: list[DimPerson]) -> list[DimPersonRole]:
    out: list[DimPersonRole] = []
    seen: set[tuple[str, int]] = set()
    for person in people:
        if person.role_id is None or int(person.role_id) <= 0:
            continue
        key = (person.person_id, int(person.role_id))
        if key in seen:
            continue
        seen.add(key)
        out.append(DimPersonRole(person_id=person.person_id, role_id=int(person.role_id)))
    return out


async def _ensure_person_extra_table(session: AsyncSession) -> None:
    try:
        await session.execute(select(DimPersonExtra.person_id).limit(1))
    except SQLAlchemyError as exc:
        message = str(exc).lower()
        missing_tokens = ("doesn't exist", "no such table", "unknown table", "undefined table")
        if "dim_person_extra" in message and any(token in message for token in missing_tokens):
            raise HTTPException(
                status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
                detail=(
                    "dim_person_extra table is missing. Apply migration "
                    "`backend/migrations/0039_dim_person_extra.sql`."
                ),
            ) from exc
        raise


async def _upsert_person_extra(
    *,
    session: AsyncSession,
    person: DimPerson,
    normalized: dict[str, Any],
    batch_hash: str | None,
    dry_run: bool,
) -> None:
    if dry_run:
        return
    now = datetime.utcnow()
    record = _build_person_extra_record(
        person=person,
        normalized=normalized,
        batch_hash=batch_hash,
        now=now,
    )
    if record is None:
        return

    existing = await session.get(DimPersonExtra, person.person_id)
    if existing is None:
        session.add(record)
        return

    if (
        existing.payload_hash == record.payload_hash
        and (existing.extra_payload_json or "") == (record.extra_payload_json or "")
        and (existing.raw_payload_json or "") == (record.raw_payload_json or "")
    ):
        if batch_hash and existing.last_seen_batch_hash != batch_hash:
            existing.last_seen_batch_hash = batch_hash
            existing.updated_at = now
        return

    existing.person_code = record.person_code
    existing.email = record.email
    existing.raw_payload_json = record.raw_payload_json
    existing.extra_payload_json = record.extra_payload_json
    existing.payload_hash = record.payload_hash
    existing.last_seen_batch_hash = record.last_seen_batch_hash
    existing.updated_at = now


def _build_person_extra_record(
    *,
    person: DimPerson,
    normalized: dict[str, Any],
    batch_hash: str | None,
    now: datetime,
) -> DimPersonExtra | None:
    raw_payload, extra_payload = _serialize_row_payloads(normalized)
    raw_json = json.dumps(raw_payload, ensure_ascii=True, sort_keys=True, separators=(",", ":"))
    extra_json = json.dumps(extra_payload, ensure_ascii=True, sort_keys=True, separators=(",", ":"))
    payload_hash = hashlib.sha256(extra_json.encode("utf-8")).hexdigest()
    return DimPersonExtra(
        person_id=person.person_id,
        person_code=person.person_code,
        email=_normalize_email(person.email) or None,
        raw_payload_json=raw_json,
        extra_payload_json=extra_json,
        payload_hash=payload_hash,
        last_seen_batch_hash=batch_hash,
        created_at=now,
        updated_at=now,
    )


def _serialize_row_payloads(normalized: dict[str, Any]) -> tuple[dict[str, Any], dict[str, Any]]:
    raw_payload: dict[str, Any] = {}
    extra_payload: dict[str, Any] = {}
    for key, value in normalized.items():
        serialized = _to_json_compatible(value)
        raw_payload[key] = serialized
        if key not in _PERSON_CORE_FIELDS:
            extra_payload[key] = serialized
    return raw_payload, extra_payload


def _parse_extra_payload_json(payload_text: str | None) -> dict[str, Any]:
    if not (payload_text or "").strip():
        return {}
    try:
        parsed = json.loads(payload_text)
    except json.JSONDecodeError:
        return {}
    return parsed if isinstance(parsed, dict) else {}


def _external_employee_code_from_raw_payload(payload: dict[str, Any]) -> str:
    direct = _normalize_text(payload.get("external_employee_code")) or _normalize_text(payload.get("employee_number"))
    if direct:
        return direct
    source = _external_identity_source_from_payload(payload)
    if source == _DEFAULT_EXTERNAL_IDENTITY_SOURCE:
        return _normalize_text(payload.get("person_code"))
    return ""


def _to_json_compatible(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (int, float, bool)):
        return value
    text = _normalize_text(value)
    return text if text != "" else None


def _load_bulk_rows(filename: str, raw: bytes) -> list[_BulkRow]:
    ext = (filename.rsplit(".", 1)[-1].lower() if "." in filename else "").strip()
    if ext in {"xlsx", "xlsm"}:
        return _read_xlsx_rows(raw)
    if ext not in {"csv", ""}:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Unsupported file type. Use .csv or .xlsx.",
        )
    return _read_csv_rows(raw)


def _read_csv_rows(raw: bytes) -> list[_BulkRow]:
    decoded = raw.decode("utf-8-sig", errors="replace")
    reader = csv.DictReader(io.StringIO(decoded))
    out: list[_BulkRow] = []
    for row_number, row in enumerate(reader, start=2):
        normalized = _normalize_row(row)
        if _is_footer_row(normalized):
            continue
        out.append(_BulkRow(row_number=row_number, normalized=normalized))
    return out


def _read_xlsx_rows(raw: bytes) -> list[_BulkRow]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - dependency guard
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="XLSX upload requires openpyxl. Install it on backend and retry.",
        ) from exc

    wb = load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
    sheet = wb[wb.sheetnames[0]]
    all_rows = list(sheet.iter_rows(values_only=True))
    if not all_rows:
        return []

    header_idx = _detect_header_row(all_rows)
    header_cells = all_rows[header_idx]
    headers = [_normalize_header_key(cell) for cell in header_cells]
    out: list[_BulkRow] = []
    for row_number, values in enumerate(all_rows[header_idx + 1 :], start=header_idx + 2):
        row: dict[str, Any] = {}
        non_empty_count = 0
        for idx, key in enumerate(headers):
            if not key or idx >= len(values):
                continue
            value = values[idx]
            if isinstance(value, str):
                value = value.strip()
            row[key] = value
            if value is not None and (not isinstance(value, str) or value.strip()):
                non_empty_count += 1
        if non_empty_count == 0:
            continue
        normalized = _normalize_row(row)
        if _is_footer_row(normalized):
            continue
        out.append(_BulkRow(row_number=row_number, normalized=normalized))
    return out


def _detect_header_row(all_rows: list[tuple[Any, ...]]) -> int:
    best_idx = 0
    best_score = -1
    scan_limit = min(len(all_rows), 15)
    indicator_headers = {"employee_number", "work_email", "first_name", "person_code", "email"}
    for idx in range(scan_limit):
        normalized_headers = {
            _normalize_header_key(value)
            for value in all_rows[idx]
            if value is not None and str(value).strip()
        }
        if not normalized_headers:
            continue
        score = len(normalized_headers.intersection(indicator_headers))
        if score > best_score:
            best_idx = idx
            best_score = score
        if score >= 2:
            return idx
    return best_idx


def _normalize_header_key(value: Any) -> str:
    text = str(value or "").strip().lower()
    return re.sub(r"[^a-z0-9]+", "_", text).strip("_")


def _dedupe_bulk_rows(rows: list[_BulkRow]) -> list[_BulkRow]:
    deduped: list[_BulkRow] = []
    key_to_idx: dict[str, int] = {}
    for row in rows:
        key = _bulk_row_key(row)
        if key in key_to_idx:
            deduped[key_to_idx[key]] = row
            continue
        key_to_idx[key] = len(deduped)
        deduped.append(row)
    return deduped


def _bulk_row_key(row: _BulkRow) -> str:
    external_employee_code = _external_employee_code(row.normalized)
    if external_employee_code:
        return f"external:{_external_identity_key(_external_identity_source_from_payload(row.normalized), external_employee_code)}"
    person_code = _incoming_internal_person_code(row.normalized)
    email = _normalize_email(row.normalized.get("email"))
    person_id = _normalize_text(row.normalized.get("person_id"))
    if person_code:
        return f"code:{person_code.lower()}"
    if email:
        return f"email:{email}"
    if person_id:
        return f"id:{person_id}"
    return f"row:{row.row_number}"


async def _load_existing_people(session: AsyncSession, rows: list[_BulkRow]) -> list[DimPerson]:
    del rows
    rows_db = await session.execute(select(DimPerson))
    return rows_db.scalars().all()


def _build_people_maps(
    people: list[DimPerson],
) -> tuple[dict[str, list[DimPerson]], dict[str, list[DimPerson]], dict[str, list[DimPerson]]]:
    by_person_id: dict[str, list[DimPerson]] = {}
    by_person_code: dict[str, list[DimPerson]] = {}
    by_email: dict[str, list[DimPerson]] = {}
    for person in people:
        by_person_id.setdefault((person.person_id or "").strip(), []).append(person)
        code = (person.person_code or "").strip()
        if code:
            by_person_code.setdefault(code.lower(), []).append(person)
        email = _normalize_email(person.email)
        if email:
            by_email.setdefault(email, []).append(person)
    return by_person_id, by_person_code, by_email


async def _load_external_employee_identity_map(
    *,
    session: AsyncSession,
    by_person_id: dict[str, list[DimPerson]],
    warnings: list[BulkUploadWarning],
) -> dict[str, list[DimPerson]]:
    rows = (
        await session.execute(
            select(
                DimPersonExtra.person_id,
                DimPersonExtra.raw_payload_json,
            )
        )
    ).all()
    by_external_employee_code: dict[str, list[DimPerson]] = {}
    for row in rows:
        person_matches = by_person_id.get(_normalize_text(row.person_id), [])
        if len(person_matches) != 1:
            continue
        payload = _parse_extra_payload_json(row.raw_payload_json)
        if not payload:
            continue
        external_employee_code = _external_employee_code_from_raw_payload(payload)
        if not external_employee_code:
            continue
        key = _external_identity_key(_external_identity_source_from_payload(payload), external_employee_code)
        by_external_employee_code.setdefault(key, []).append(person_matches[0])

    for key, matches in by_external_employee_code.items():
        unique_person_ids = {person.person_id for person in matches}
        if len(unique_person_ids) > 1:
            warnings.append(
                BulkUploadWarning(
                    message=(
                        f"Existing external employee mapping '{key}' points to multiple people. "
                        "Resolve the duplicate mapping before relying on automatic reconciliation."
                    ),
                )
            )
    return by_external_employee_code


def _resolve_existing_person(
    *,
    row: _BulkRow,
    by_person_id: dict[str, list[DimPerson]],
    by_person_code: dict[str, list[DimPerson]],
    by_external_employee_code: dict[str, list[DimPerson]],
    by_email: dict[str, list[DimPerson]],
    warnings: list[BulkUploadWarning],
) -> DimPerson | None:
    normalized = row.normalized
    person_id = _normalize_text(normalized.get("person_id"))
    person_code = _incoming_internal_person_code(normalized)
    external_employee_code = _external_employee_code(normalized)
    email = _normalize_email(normalized.get("email"))

    match_by_id = _single_person_match(by_person_id, person_id, "person_id") if person_id else None
    match_by_external = (
        _single_person_match(
            by_external_employee_code,
            _external_identity_key(_external_identity_source_from_payload(normalized), external_employee_code),
            "external_employee_code",
        )
        if external_employee_code
        else None
    )
    match_by_code = _single_person_match(by_person_code, person_code.lower() if person_code else "", "person_code") if person_code else None
    match_by_email = _single_person_match(by_email, email, "email") if email else None

    anchor = match_by_external or match_by_code or match_by_email or match_by_id
    if anchor is None:
        return None

    if match_by_external and match_by_email and match_by_external.person_id != match_by_email.person_id:
        warnings.append(
            BulkUploadWarning(
                row=row.row_number,
                message=(
                    f"Email '{email}' belongs to another record. External employee code match was used; email was ignored."
                ),
                person_id=anchor.person_id,
                person_code=person_code or external_employee_code,
                email=email,
            )
        )
        normalized.pop("email", None)
    elif match_by_external and match_by_code and match_by_external.person_id != match_by_code.person_id:
        warnings.append(
            BulkUploadWarning(
                row=row.row_number,
                message="External employee code conflicts with incoming internal person_code. External mapping was kept.",
                person_id=anchor.person_id,
                person_code=person_code or external_employee_code,
                email=email,
            )
        )
        normalized.pop("person_code", None)
    elif match_by_code and match_by_email and match_by_code.person_id != match_by_email.person_id:
        # Canonical internal business key wins: preserve anchor resolved by person_code and avoid email hijack.
        warnings.append(
            BulkUploadWarning(
                row=row.row_number,
                message=(
                    f"Email '{email}' belongs to another record. Person code match was used; email was ignored."
                ),
                person_id=anchor.person_id,
                person_code=person_code,
                email=email,
            )
        )
        normalized.pop("email", None)
    elif match_by_id and anchor.person_id != match_by_id.person_id:
        warnings.append(
            BulkUploadWarning(
                row=row.row_number,
                message="Incoming person_id conflicts with resolved record. Existing mapping was kept.",
                person_id=anchor.person_id,
                person_code=person_code,
                email=email,
            )
        )
    return anchor


def _single_person_match(
    mapping: dict[str, list[DimPerson]],
    key: str,
    label: str,
) -> DimPerson | None:
    if not key:
        return None
    matches = mapping.get(key, [])
    if not matches:
        return None
    if len(matches) > 1:
        raise ValueError(f"Multiple records found for {label} '{key}'. Resolve duplicates first.")
    return matches[0]


def _build_safe_update_payload(existing: DimPerson, updates: dict[str, Any]) -> dict[str, Any]:
    patch: dict[str, Any] = {}
    for key, value in updates.items():
        if key in {"person_id", "created_at"}:
            continue
        if key == "updated_at":
            patch[key] = value
            continue
        if value is None and key not in _UPDATABLE_TO_NONE_FIELDS:
            continue
        if key == "email":
            value = _normalize_email(value)
        current = getattr(existing, key)
        if current != value:
            patch[key] = value
    return patch


def _preserve_existing_person_code(
    *,
    existing: DimPerson,
    updates: dict[str, Any],
) -> bool:
    incoming_code = _normalize_text(updates.get("person_code"))
    existing_code = _normalize_text(existing.person_code)
    if not incoming_code or not existing_code or incoming_code == existing_code:
        return False

    updates["person_code"] = existing_code
    return True


def _assert_no_update_conflicts(
    existing: DimPerson,
    patch: dict[str, Any],
    by_person_code: dict[str, list[DimPerson]],
    by_email: dict[str, list[DimPerson]],
) -> None:
    new_code = _normalize_text(patch.get("person_code"))
    if new_code:
        for person in by_person_code.get(new_code.lower(), []):
            if person.person_id != existing.person_id:
                raise ValueError(f"person_code '{new_code}' is already used by another person.")
    new_email = _normalize_email(patch.get("email"))
    if new_email:
        for person in by_email.get(new_email, []):
            if person.person_id != existing.person_id:
                raise ValueError(f"email '{new_email}' is already used by another person.")


def _assert_no_create_conflicts(
    payload: dict[str, Any],
    by_person_code: dict[str, list[DimPerson]],
    by_email: dict[str, list[DimPerson]],
) -> None:
    person_code = _normalize_text(payload.get("person_code"))
    email = _normalize_email(payload.get("email"))
    if person_code and by_person_code.get(person_code.lower()):
        raise ValueError(f"person_code '{person_code}' already exists.")
    if email and by_email.get(email):
        raise ValueError(f"email '{email}' already exists.")


def _refresh_person_maps(
    *,
    person: DimPerson,
    by_person_id: dict[str, list[DimPerson]],
    by_person_code: dict[str, list[DimPerson]],
    by_email: dict[str, list[DimPerson]],
) -> None:
    for mapping in (by_person_code, by_email):
        for key in list(mapping.keys()):
            mapping[key] = [item for item in mapping[key] if item.person_id != person.person_id]
            if not mapping[key]:
                mapping.pop(key, None)
    by_person_id[(person.person_id or "").strip()] = [person]
    person_code = _normalize_text(person.person_code)
    if person_code:
        by_person_code.setdefault(person_code.lower(), []).append(person)
    email = _normalize_email(person.email)
    if email:
        by_email.setdefault(email, []).append(person)


def _refresh_external_employee_identity_map(
    *,
    person: DimPerson,
    normalized: dict[str, Any],
    by_external_employee_code: dict[str, list[DimPerson]],
) -> None:
    external_employee_code = _external_employee_code(normalized)
    if not external_employee_code:
        return
    key = _external_identity_key(_external_identity_source_from_payload(normalized), external_employee_code)
    for existing_key in list(by_external_employee_code.keys()):
        by_external_employee_code[existing_key] = [
            item for item in by_external_employee_code[existing_key] if item.person_id != person.person_id
        ]
        if not by_external_employee_code[existing_key]:
            by_external_employee_code.pop(existing_key, None)
    by_external_employee_code.setdefault(key, []).append(person)


def _resolve_manager_id(
    manager_value: Any,
    by_person_id: dict[str, list[DimPerson]],
    by_person_code: dict[str, list[DimPerson]],
) -> str | None:
    manager_token = _normalize_text(manager_value)
    if not manager_token:
        return None
    direct = by_person_id.get(manager_token, [])
    if len(direct) == 1:
        return direct[0].person_id
    by_code = by_person_code.get(manager_token.lower(), [])
    if len(by_code) == 1:
        return by_code[0].person_id
    return None


def _requested_manager_reference(normalized: dict[str, Any]) -> str:
    return _external_manager_code(normalized) or _normalize_text(normalized.get("manager_id"))


def _resolve_manager_reference(
    *,
    normalized: dict[str, Any],
    by_person_id: dict[str, list[DimPerson]],
    by_person_code: dict[str, list[DimPerson]],
    by_external_employee_code: dict[str, list[DimPerson]],
) -> str | None:
    external_manager_code = _external_manager_code(normalized)
    direct_manager_value = _normalize_text(normalized.get("manager_id"))
    if external_manager_code:
        manager = _single_person_match(
            by_external_employee_code,
            _external_identity_key(_external_identity_source_from_payload(normalized), external_manager_code),
            "external_manager_code",
        )
        if manager is not None:
            return manager.person_id
        if direct_manager_value and direct_manager_value != external_manager_code:
            return _resolve_manager_id(direct_manager_value, by_person_id, by_person_code)
        return None
    return _resolve_manager_id(direct_manager_value, by_person_id, by_person_code)


@router.get("/{person_id}", response_model=PlatformPersonOut)
async def get_person(
    person_id: str,
    session: AsyncSession = Depends(get_platform_session),
    _user: UserContext = Depends(require_superadmin()),
):
    row = (
        await session.execute(
            select(DimPerson, DimRole.role_code, DimRole.role_name)
            .select_from(DimPerson)
            .outerjoin(DimRole, DimRole.role_id == DimPerson.role_id)
            .where(DimPerson.person_id == person_id)
        )
    ).first()
    if not row:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Person not found")
    person, role_code, role_name = row
    return _person_out(person, role_code, role_name)


@router.post("", response_model=PlatformPersonOut, status_code=status.HTTP_201_CREATED)
async def create_person(
    payload: PlatformPersonCreate,
    session: AsyncSession = Depends(get_platform_session),
    _user: UserContext = Depends(require_superadmin()),
):
    person_id = (payload.person_id or "").strip() or None
    if person_id and person_id.isdigit():
        person_id = None
    if not person_id:
        person_id = await _generate_person_id(session, payload.first_name, payload.last_name)
    existing = await session.get(DimPerson, person_id)
    if existing:
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="person_id already exists")
    person = DimPerson(**payload.model_dump(exclude={"person_id"}), person_id=person_id)
    if person.created_at is None:
        person.created_at = datetime.utcnow()
    if person.updated_at is None:
        person.updated_at = datetime.utcnow()
    session.add(person)
    await session.commit()
    await session.refresh(person)
    role_code, role_name = await _role_meta(session, person.role_id)
    return _person_out(person, role_code, role_name)


@router.patch("/{person_id}", response_model=PlatformPersonOut)
async def update_person(
    person_id: str,
    payload: PlatformPersonUpdate,
    session: AsyncSession = Depends(get_platform_session),
    _user: UserContext = Depends(require_superadmin()),
):
    person = await session.get(DimPerson, person_id)
    if not person:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Person not found")
    updates = _clean_update_payload(payload)
    for key, value in updates.items():
        setattr(person, key, value)
    if "updated_at" not in updates:
        person.updated_at = datetime.utcnow()
    await session.commit()
    await session.refresh(person)
    role_code, role_name = await _role_meta(session, person.role_id)
    return _person_out(person, role_code, role_name)


@router.delete("/{person_id}", status_code=status.HTTP_200_OK)
async def delete_person(
    person_id: str,
    session: AsyncSession = Depends(get_platform_session),
    _user: UserContext = Depends(require_superadmin()),
):
    person = await session.get(DimPerson, person_id)
    if not person:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Person not found")
    person.is_deleted = 1
    person.updated_at = datetime.utcnow()
    await session.commit()
    return {"person_id": person_id, "is_deleted": 1}


async def _role_meta(session: AsyncSession, role_id: int | None) -> tuple[str | None, str | None]:
    if role_id is None:
        return None, None
    row = (
        await session.execute(
            select(DimRole.role_code, DimRole.role_name).where(DimRole.role_id == role_id)
        )
    ).first()
    if not row:
        return None, None
    return row[0], row[1]


def _person_out(person: DimPerson, role_code: str | None, role_name: str | None) -> PlatformPersonOut:
    return PlatformPersonOut(
        person_id=person.person_id,
        person_code=person.person_code,
        personal_id=person.personal_id,
        first_name=person.first_name,
        last_name=person.last_name,
        email=person.email,
        mobile_number=person.mobile_number,
        role_id=person.role_id,
        grade_id=person.grade_id,
        department_id=person.department_id,
        manager_id=person.manager_id,
        employment_type=person.employment_type,
        join_date=person.join_date,
        exit_date=person.exit_date,
        status=person.status,
        is_deleted=person.is_deleted,
        created_at=person.created_at,
        updated_at=person.updated_at,
        source_system=person.source_system,
        source_candidate_id=person.source_candidate_id,
        source_candidate_code=person.source_candidate_code,
        full_name=person.full_name,
        display_name=person.display_name,
        role_code=role_code,
        role_name=role_name,
    )


def _clean_update_payload(payload: PlatformPersonUpdate) -> dict:
    updates = payload.model_dump(exclude_unset=True)
    return _coerce_payload(updates)


def _normalize_text(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _normalize_email(value: object) -> str:
    text = _normalize_text(value)
    return text.lower() if text else ""


def _is_nullish_text(value: str) -> bool:
    normalized = re.sub(r"[^a-z0-9]+", "_", value.strip().lower()).strip("_")
    return normalized in _NULLISH_TEXT


def _is_footer_row(normalized: dict[str, Any]) -> bool:
    person_code = _normalize_text(normalized.get("person_code")).lower()
    email = _normalize_email(normalized.get("email"))
    first_name = _normalize_text(normalized.get("first_name"))
    if person_code.startswith("generated on"):
        return True
    if not person_code and not email and not first_name:
        return True
    return False


def _apply_name_defaults(payload: dict[str, Any]) -> dict[str, Any]:
    out = dict(payload)
    first_name = _normalize_text(out.get("first_name"))
    last_name = _normalize_text(out.get("last_name"))
    full_name = _normalize_text(out.get("full_name"))
    display_name = _normalize_text(out.get("display_name"))
    if not full_name:
        composed = " ".join(part for part in [first_name, last_name] if part).strip()
        if composed:
            out["full_name"] = composed
            full_name = composed
    if not display_name and full_name:
        out["display_name"] = full_name
    return out


def _apply_status_defaults(payload: dict[str, Any]) -> dict[str, Any]:
    out = dict(payload)
    status_text = _normalize_text(out.get("status"))
    if status_text and "is_deleted" not in out:
        out["is_deleted"] = _derive_is_deleted_from_status(status_text)
    return out


def _apply_external_identity_defaults(normalized: dict[str, Any]) -> None:
    external_employee_code = _external_employee_code(normalized)
    if not external_employee_code:
        return
    normalized["external_employee_code"] = external_employee_code
    external_manager_code = _external_manager_code(normalized)
    if external_manager_code:
        normalized["external_manager_code"] = external_manager_code
    if not _normalize_text(normalized.get("external_identity_source")):
        normalized["external_identity_source"] = _normalize_text(normalized.get("source_system")) or _DEFAULT_EXTERNAL_IDENTITY_SOURCE


def _normalize_row(row: dict[str, Any]) -> dict[str, Any]:
    normalized: dict[str, Any] = {}
    for key, value in row.items():
        if key is None:
            continue
        norm_key = _normalize_header_key(key)
        if norm_key == "employee_number":
            normalized["external_employee_code"] = value.strip() if isinstance(value, str) else value
        if norm_key == "reporting_manager_employee_number":
            normalized["external_manager_code"] = value.strip() if isinstance(value, str) else value
        mapped_key = _HEADER_ALIASES.get(norm_key, norm_key)
        mapped_value: Any = value.strip() if isinstance(value, str) else value
        if mapped_key in {"email"} and isinstance(mapped_value, str):
            mapped_value = mapped_value.lower()
        if mapped_key in normalized and normalized[mapped_key] not in {None, ""} and mapped_value in {None, ""}:
            continue
        normalized[mapped_key] = mapped_value
    return normalized


def _apply_row_updates(normalized: dict[str, Any]) -> dict:
    updates = {k: normalized.get(k) for k in _PERSON_CORE_FIELDS if k in normalized}
    coerced = _coerce_payload(updates)
    if coerced.get("person_id") is None:
        coerced.pop("person_id", None)
    return coerced


def _coerce_payload(payload: dict) -> dict:
    out: dict = {}
    for key, value in payload.items():
        if isinstance(value, str):
            trimmed = value.strip()
            if trimmed == "" or _is_nullish_text(trimmed):
                value = None
            else:
                value = trimmed
        if key in {"role_id", "grade_id", "department_id", "source_candidate_id"}:
            out[key] = _parse_int(value, key)
        elif key in {"is_deleted"}:
            out[key] = _parse_bool_int(value)
        elif key in {"join_date", "exit_date", "date_of_birth", "marriage_date"}:
            out[key] = _parse_date(value, key)
        elif key in {"created_at", "updated_at"}:
            out[key] = _parse_datetime(value, key)
        else:
            if value is None:
                out[key] = None
            elif isinstance(value, str):
                out[key] = value
            else:
                out[key] = str(value)
    status = (out.get("status") or "").strip().lower() if isinstance(out.get("status"), str) else out.get("status")
    if status == "relieved":
        out["role_id"] = 0
    return out


def _parse_int(value: object, field: str) -> int | None:
    if value is None:
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        if value.is_integer():
            return int(value)
        raise ValueError(f"Invalid {field}")
    if isinstance(value, str):
        cleaned = value.strip()
        if cleaned.isdigit():
            return int(cleaned)
        try:
            as_float = float(cleaned)
        except ValueError:
            raise ValueError(f"Invalid {field}") from None
        if as_float.is_integer():
            return int(as_float)
        raise ValueError(f"Invalid {field}")
    raise ValueError(f"Invalid {field}")


def _parse_bool_int(value: object) -> int | None:
    if value is None:
        return None
    if isinstance(value, int):
        return 1 if value else 0
    if isinstance(value, float):
        return 1 if value else 0
    if isinstance(value, str):
        raw = value.strip().lower()
        if raw in {"1", "true", "yes"}:
            return 1
        if raw in {"0", "false", "no"}:
            return 0
    raise ValueError("Invalid is_deleted")


def _parse_date(value: object, field: str) -> date | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date) and not isinstance(value, datetime):
        return value
    if isinstance(value, str):
        raw = value.strip()
        if not raw:
            return None
        try:
            return date.fromisoformat(raw)
        except ValueError:
            pass
        for fmt in (
            "%d-%b-%y",
            "%d-%b-%Y",
            "%d-%m-%y",
            "%d-%m-%Y",
            "%Y/%m/%d",
        ):
            try:
                return datetime.strptime(raw, fmt).date()
            except ValueError:
                continue
        slash_match = re.fullmatch(r"(\d{1,2})/(\d{1,2})/(\d{2,4})", raw)
        if slash_match:
            first = int(slash_match.group(1))
            second = int(slash_match.group(2))
            year = slash_match.group(3)
            if first > 12:
                fmts = ("%d/%m/%Y", "%d/%m/%y")
            elif second > 12:
                fmts = ("%m/%d/%Y", "%m/%d/%y")
            else:
                # Ambiguous `x/y/yyyy`: prefer month/day for Excel exports used in this flow.
                fmts = ("%m/%d/%Y", "%m/%d/%y", "%d/%m/%Y", "%d/%m/%y")
            for fmt in fmts:
                if len(year) == 2 and not fmt.endswith("%y"):
                    continue
                if len(year) == 4 and not fmt.endswith("%Y"):
                    continue
                try:
                    return datetime.strptime(raw, fmt).date()
                except ValueError:
                    continue
        try:
            return _parse_datetime(raw, field).date()
        except ValueError as exc:
            raise ValueError(f"Invalid {field}: {raw}") from exc
    raise ValueError(f"Invalid {field}")


def _parse_datetime(value: object, field: str) -> datetime | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value
    if isinstance(value, str):
        raw = value.strip()
        if not raw:
            return None
        normalized = raw.replace("Z", "+00:00")
        try:
            return datetime.fromisoformat(normalized)
        except ValueError:
            pass
        for fmt in (
            "%m/%d/%YT%I:%M:%S %p",
            "%m/%d/%Y %I:%M:%S %p",
            "%m/%d/%YT%H:%M:%S",
            "%m/%d/%Y %H:%M:%S",
            "%m/%d/%Y",
        ):
            try:
                return datetime.strptime(raw, fmt)
            except ValueError:
                continue
        raise ValueError(f"Invalid {field}: {raw}")
    raise ValueError(f"Invalid {field}")


def _build_person_prefix(first_name: str | None, last_name: str | None) -> str:
    first = (first_name or "").strip()
    last = (last_name or "").strip()
    if first and last:
        return f"{_first_letter(first)}{_first_letter(last)}"
    if first:
        return f"{_first_letter(first)}{_second_letter_or(first, 'X')}"
    return "XX"


def _first_letter(value: str) -> str:
    for ch in value:
        if ch.isalpha():
            return ch.upper()
    return "X"


def _second_letter_or(value: str, fallback: str) -> str:
    letters = [ch.upper() for ch in value if ch.isalpha()]
    if len(letters) >= 2:
        return letters[1]
    return fallback


async def _lock_person_code_counters(session: AsyncSession) -> dict[str, _PersonCodeCounterState]:
    rows = (
        await session.execute(
            text(
                """
                SELECT `code_type`, `next_seq`, `max_seq`
                FROM `dim_person_code_counter`
                FOR UPDATE
                """
            )
        )
    ).all()
    states: dict[str, _PersonCodeCounterState] = {}
    for row in rows:
        code_type = _normalize_text(row.code_type).upper()
        next_seq = int(row.next_seq or 0)
        states[code_type] = _PersonCodeCounterState(
            code_type=code_type,
            next_seq=next_seq,
            original_next_seq=next_seq,
            max_seq=int(row.max_seq) if row.max_seq is not None else None,
        )
    return states


async def _normalize_bulk_employment_type(
    *,
    session: AsyncSession,
    value: str | None,
    job_title: str | None,
    email: str | None,
) -> str:
    normalized = (
        await session.execute(
            text(
                """
                SELECT fn_dim_person_normalize_employment_type(
                  :employment_type,
                  :job_title,
                  :email
                ) AS normalized_employment_type
                """
            ),
            {
                "employment_type": _normalize_text(value) or None,
                "job_title": _normalize_text(job_title) or None,
                "email": _normalize_email(email) or None,
            },
        )
    ).scalar_one_or_none()
    return str(normalized or "Permanent").strip() or "Permanent"


async def _bulk_person_code_type(
    *,
    session: AsyncSession,
    employment_type: str,
) -> str:
    code_type = (
        await session.execute(
            text("SELECT fn_dim_person_code_type(:principal_flag, :employment_type) AS code_type"),
            {"principal_flag": 0, "employment_type": employment_type},
        )
    ).scalar_one_or_none()
    if not code_type:
        raise ValueError("Could not resolve person_code series for upload row.")
    return str(code_type).strip().upper()


async def _bulk_person_code_is_valid(
    *,
    session: AsyncSession,
    code_type: str,
    person_code: str | None,
) -> bool:
    candidate = _normalize_text(person_code).upper()
    if not candidate:
        return False
    is_valid = (
        await session.execute(
            text("SELECT fn_dim_person_code_is_valid(:code_type, :person_code) AS is_valid"),
            {"code_type": code_type, "person_code": candidate},
        )
    ).scalar_one_or_none()
    return bool(int(is_valid or 0))


async def _format_bulk_person_code(
    *,
    session: AsyncSession,
    code_type: str,
    seq: int,
) -> str:
    candidate = (
        await session.execute(
            text("SELECT fn_dim_person_format_code(:code_type, :seq) AS candidate_code"),
            {"code_type": code_type, "seq": seq},
        )
    ).scalar_one_or_none()
    if not candidate:
        raise ValueError("Could not format person_code for upload row.")
    return str(candidate).strip().upper()


async def _resolve_bulk_person_code(
    *,
    session: AsyncSession,
    normalized: dict[str, Any],
    payload: dict[str, Any],
    by_person_code: dict[str, list[DimPerson]],
    counters: dict[str, _PersonCodeCounterState],
) -> tuple[str, str]:
    normalized_employment_type = await _normalize_bulk_employment_type(
        session=session,
        value=_normalize_text(payload.get("employment_type")) or None,
        job_title=_normalize_text(payload.get("job_title")) or None,
        email=_normalize_email(payload.get("email")) or None,
    )
    code_type = await _bulk_person_code_type(session=session, employment_type=normalized_employment_type)
    explicit_internal_code = _incoming_internal_person_code(normalized)
    if explicit_internal_code:
        explicit_internal_code = explicit_internal_code.upper()
        if await _bulk_person_code_is_valid(
            session=session,
            code_type=code_type,
            person_code=explicit_internal_code,
        ) and not by_person_code.get(explicit_internal_code.lower()):
            return explicit_internal_code, normalized_employment_type

    state = counters.get(code_type)
    if state is None:
        raise ValueError(f"Person code counter is not configured for code type '{code_type}'.")

    seq = state.next_seq
    while True:
        if state.max_seq is not None and seq > state.max_seq:
            raise ValueError(f"No remaining person_code sequence for code type '{code_type}'.")
        candidate = await _format_bulk_person_code(session=session, code_type=code_type, seq=seq)
        seq += 1
        if by_person_code.get(candidate.lower()):
            continue
        state.next_seq = seq
        return candidate, normalized_employment_type


async def _save_person_code_counters(
    session: AsyncSession,
    counters: dict[str, _PersonCodeCounterState],
) -> None:
    for state in counters.values():
        if state.next_seq == state.original_next_seq:
            continue
        await session.execute(
            text(
                """
                UPDATE `dim_person_code_counter`
                SET `next_seq` = :next_seq,
                    `updated_at` = NOW()
                WHERE `code_type` = :code_type
                """
            ),
            {"code_type": state.code_type, "next_seq": state.next_seq},
        )


async def _generate_person_id(session: AsyncSession, first_name: str | None, last_name: str | None) -> str:
    next_seq = await _ensure_next_seq(session, None)
    return _format_person_id(first_name, last_name, next_seq)


async def _ensure_next_seq(session: AsyncSession, current: int | None) -> int:
    if current is not None:
        return current
    pattern = r"^[A-Z]{2}_[0-9]+$"
    max_row = (
        await session.execute(
            select(
                func.max(
                    func.cast(
                        func.substr(DimPerson.person_id, func.instr(DimPerson.person_id, "_") + 1, 10),
                        Integer,
                    )
                )
            ).where(DimPerson.person_id.op("REGEXP")(pattern))
        )
    ).first()
    max_value = max_row[0] if max_row and max_row[0] is not None else 0
    return int(max_value) + 1


def _format_person_id(first_name: str | None, last_name: str | None, seq: int) -> str:
    prefix = _build_person_prefix(first_name, last_name)
    return f"{prefix}_{str(seq).zfill(3)}"


def _consume_next_seq(prefix: str, seq: int, used_ids: set[str]) -> int:
    while True:
        candidate = f"{prefix}_{str(seq).zfill(3)}"
        if candidate not in used_ids:
            used_ids.add(candidate)
            return seq
        seq += 1
